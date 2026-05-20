from __future__ import annotations

import json
import multiprocessing
import os
import platform
import re
import shutil
import signal
import stat
import subprocess
import sys
import tempfile
import time
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

import openpyxl

from browser_startup import BrowserStartupError, BrowserStartupInfo, prepare_chromedriver
from models import ScrapeItem, TaskRequest, TaskResult
from trip_cli_scraper import TripCliScraper
from xlsx_export import export_rows

Reporter = Callable[[dict[str, Any] | str], None]
DEFAULT_VENDOR = "手动输入"
MAX_WORKERS = 5
try:
    import psutil
except Exception:
    psutil = None


def load_inputs(request: TaskRequest) -> list[str]:
    return [item.url for item in load_items(request)]


def load_items(request: TaskRequest) -> list[ScrapeItem]:
    items: list[ScrapeItem] = []
    order = 0
    for url in request.input_paths:
        if url.strip():
            items.append(ScrapeItem(url=url.strip(), vendor=DEFAULT_VENDOR, input_order=order))
            order += 1
    if request.urls_file:
        lines = Path(request.urls_file).expanduser().read_text(encoding="utf-8").splitlines()
        for row_number, line in enumerate(lines, 1):
            url = line.strip()
            if url:
                items.append(ScrapeItem(url=url, vendor=DEFAULT_VENDOR, source_row=row_number, input_order=order))
                order += 1
    if request.excel_input:
        excel_items, _warnings = load_excel_items(request.excel_input, start_order=order)
        items.extend(excel_items)
    if not items:
        raise ValueError("At least one URL is required via --input, --urls-file, or --excel-input.")
    return items


def load_excel_items(path: str, start_order: int = 0) -> tuple[list[ScrapeItem], list[str]]:
    workbook_path = Path(path).expanduser()
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook.active
    items: list[ScrapeItem] = []
    warnings: list[str] = []
    order = start_order
    try:
        max_column = worksheet.max_column or 0
        max_row = worksheet.max_row or 0
        for column in range(1, max_column + 1):
            header_value = worksheet.cell(row=1, column=column).value
            vendor = str(header_value).strip() if header_value is not None else ""
            if not vendor:
                continue
            found = 0
            for row in range(2, max_row + 1):
                cell_value = worksheet.cell(row=row, column=column).value
                if cell_value is None:
                    continue
                url = str(cell_value).strip()
                if not url:
                    continue
                lowered = url.lower()
                if "http" in lowered and "trip.com" in lowered:
                    items.append(
                        ScrapeItem(
                            url=url,
                            vendor=vendor,
                            source_column=column,
                            source_row=row,
                            input_order=order,
                        )
                    )
                    order += 1
                    found += 1
                else:
                    warnings.append(f"ignored non-Trip link at column {column}, row {row}: {url}")
            if not found:
                warnings.append(f"vendor column {column} ({vendor}) has no Trip links")
    finally:
        workbook.close()
    return items, warnings


def default_output_path(request: TaskRequest) -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = (
        f"trip_score_results_{timestamp}.xlsx"
        if request.rating_only
        else f"trip_scoreandprice_results_{timestamp}.xlsx"
    )
    return str(Path.home() / "Documents" / filename)


def default_vendor_output_dir(base_dir: Path | None = None, timestamp: str | None = None) -> str:
    timestamp = timestamp or datetime.now().strftime("%y%m%d-%H%M")
    return str((base_dir or (Path.home() / "Documents")) / "trip_vendor_results" / timestamp)


def default_session_path() -> str:
    if platform.system() == "Windows":
        config_dir = Path(os.environ.get("PROGRAMDATA", "C:\\ProgramData")) / "TripDetailFilling"
    elif platform.system() == "Darwin":
        config_dir = Path.home() / "Library" / "Preferences" / "TripDetailFilling"
    else:
        config_dir = Path.home() / ".config" / "TripDetailFilling"
    return str(config_dir / "session_data.json")


def _expand_output_path(path: str) -> str:
    return str(Path(path).expanduser())


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", value).strip(" ._")
    return cleaned or "未命名厂商"


def _vendor_groups(items: list[ScrapeItem]) -> list[dict[str, Any]]:
    groups: list[dict[str, Any]] = []
    seen: dict[str, int] = {}
    by_key: dict[tuple[str, int | None], dict[str, Any]] = {}
    for item in sorted(items, key=lambda current: current.input_order):
        key = (item.vendor or DEFAULT_VENDOR, item.source_column)
        if key not in by_key:
            vendor = item.vendor or DEFAULT_VENDOR
            seen[vendor] = seen.get(vendor, 0) + 1
            suffix = f"_列{item.source_column}" if seen[vendor] > 1 and item.source_column else ""
            file_vendor = f"{vendor}{suffix}" if suffix else vendor
            group = {"vendor": vendor, "file_vendor": file_vendor, "items": []}
            by_key[key] = group
            groups.append(group)
        by_key[key]["items"].append(item)
    return groups


def _resolve_worker_count(requested: int | None, vendor_count: int) -> int:
    if requested is not None and not 1 <= requested <= MAX_WORKERS:
        raise ValueError(f"--workers must be between 1 and {MAX_WORKERS}.")
    return requested or max(1, min(vendor_count, MAX_WORKERS))


def _vendor_output_path(group: dict[str, Any], output_dir: str, timestamp: str) -> str:
    return str(Path(output_dir).expanduser() / f"{_safe_filename(group['file_vendor'])}_{timestamp}.xlsx")


def _scrape_item_to_dict(item: ScrapeItem) -> dict[str, Any]:
    return {
        "url": item.url,
        "vendor": item.vendor,
        "source_column": item.source_column,
        "source_row": item.source_row,
        "input_order": item.input_order,
    }


def _scrape_item_from_dict(data: dict[str, Any]) -> ScrapeItem:
    return ScrapeItem(
        url=str(data.get("url") or ""),
        vendor=str(data.get("vendor") or DEFAULT_VENDOR),
        source_column=data.get("source_column"),
        source_row=data.get("source_row"),
        input_order=int(data.get("input_order") or 0),
    )


def _write_json(path: str | Path, data: Any) -> None:
    target = Path(path).expanduser()
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def _read_json(path: str | Path) -> Any:
    return json.loads(Path(path).expanduser().read_text(encoding="utf-8"))


class FileStopEvent:
    def __init__(self, path: str | Path | None) -> None:
        self.path = Path(path).expanduser() if path else None

    def is_set(self) -> bool:
        return bool(self.path and self.path.exists())


def _worker_dir_path(group: dict[str, Any], workers_dir: str | Path, index: int) -> Path:
    return Path(workers_dir).expanduser() / f"{index:02d}_{_safe_filename(group['file_vendor'])}"


def _copy_driver_to_worker(driver_path: str | None, worker_dir: str | Path) -> str | None:
    if not driver_path:
        return None
    source = Path(driver_path).expanduser()
    target_dir = Path(worker_dir).expanduser() / "driver"
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / source.name
    shutil.copy2(source, target)
    try:
        target.chmod(target.stat().st_mode | stat.S_IXUSR | stat.S_IXGRP | stat.S_IXOTH)
    except Exception:
        pass
    return str(target)


def _write_stop_request(stop_file: str | Path | None) -> None:
    if not stop_file:
        return
    target = Path(stop_file).expanduser()
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(datetime.now().isoformat(timespec="seconds"), encoding="utf-8")


def _worker_popen_kwargs() -> dict[str, Any]:
    kwargs: dict[str, Any] = {"stdout": subprocess.DEVNULL, "stderr": subprocess.DEVNULL}
    if platform.system() == "Windows":
        kwargs["creationflags"] = getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0)
    else:
        kwargs["start_new_session"] = True
    return kwargs


def _terminate_process_tree(process: subprocess.Popen) -> None:
    if process.poll() is not None:
        return
    if psutil:
        try:
            parent = psutil.Process(process.pid)
            children = parent.children(recursive=True)
            for child in children:
                try:
                    child.terminate()
                except Exception:
                    pass
            try:
                parent.terminate()
            except Exception:
                pass
            _, alive = psutil.wait_procs([parent, *children], timeout=3)
            for proc in alive:
                try:
                    proc.kill()
                except Exception:
                    pass
            return
        except Exception:
            pass
    if platform.system() != "Windows":
        try:
            os.killpg(process.pid, signal.SIGTERM)
            return
        except Exception:
            pass
    try:
        process.terminate()
    except Exception:
        pass


def _build_worker_command(
    *,
    request: TaskRequest,
    vendor: str,
    items_file: str,
    output_path: str,
    result_file: str,
    session_file: str,
) -> list[str]:
    main_path = Path(__file__).with_name("main.py")
    command = [
        sys.executable,
        str(main_path),
        "--worker-items-file",
        items_file,
        "--worker-output",
        output_path,
        "--worker-result-file",
        result_file,
        "--worker-vendor",
        vendor,
        "--session-file",
        session_file,
        "--browser",
        request.browser_options.browser,
        "--timeout",
        str(request.browser_options.timeout),
        "--browser-memory-limit-mb",
        str(request.browser_options.memory_limit_mb),
        "--browser-restart-every",
        str(request.browser_options.restart_every),
    ]
    command.append("--headless" if request.browser_options.headless else "--no-headless")
    command.append("--rating-only" if request.rating_only else "--no-rating-only")
    command.append("--real-price" if request.extract_real_price else "--no-real-price")
    command.append("--product-info" if request.extract_product_info else "--no-product-info")
    if not request.browser_options.memory_guard_enabled:
        command.append("--disable-browser-memory-guard")
    if request.browser_options.chrome_path:
        command.extend(["--worker-chrome-path", request.browser_options.chrome_path])
    if request.browser_options.chrome_version_main:
        command.extend(["--worker-chrome-version-main", str(request.browser_options.chrome_version_main)])
    if request.browser_options.driver_executable_path:
        command.extend(["--worker-driver-executable-path", request.browser_options.driver_executable_path])
    if request.browser_options.user_data_dir:
        command.extend(["--worker-profile-dir", request.browser_options.user_data_dir])
    if request.stop_file:
        command.extend(["--worker-stop-file", request.stop_file])
    return command


def _create_worker_spec(
    *,
    request: TaskRequest,
    group: dict[str, Any],
    worker_id: int,
    workers_dir: str | Path,
    output_path: str,
    base_session: str,
) -> dict[str, Any]:
    worker_dir = _worker_dir_path(group, workers_dir, worker_id)
    worker_dir.mkdir(parents=True, exist_ok=True)
    items_file = worker_dir / "items.json"
    result_file = worker_dir / "result.json"
    session_file = worker_dir / "session.json"
    profile_dir = worker_dir / "profile"
    stop_file = worker_dir / "stop.request"
    output_path = _expand_output_path(output_path)

    if stop_file.exists():
        stop_file.unlink()
    _write_json(items_file, [_scrape_item_to_dict(item) for item in group["items"]])
    base_path = Path(base_session).expanduser()
    session_file.parent.mkdir(parents=True, exist_ok=True)
    if base_path.exists():
        shutil.copy2(base_path, session_file)
    driver_path = _copy_driver_to_worker(request.browser_options.driver_executable_path, worker_dir)

    worker_request = _worker_request(request, output_path, str(session_file))
    worker_request.stop_file = str(stop_file)
    worker_request.browser_options.driver_executable_path = driver_path
    worker_request.browser_options.user_data_dir = str(profile_dir)
    command = _build_worker_command(
        request=worker_request,
        vendor=group["vendor"],
        items_file=str(items_file),
        output_path=output_path,
        result_file=str(result_file),
        session_file=str(session_file),
    )
    return {
        "command": command,
        "group": group,
        "worker_id": worker_id,
        "worker_dir": str(worker_dir),
        "items_file": str(items_file),
        "result_file": str(result_file),
        "session_file": str(session_file),
        "profile_dir": str(profile_dir),
        "stop_file": str(stop_file),
        "driver_executable_path": driver_path,
        "output_path": output_path,
    }


def run_vendor_worker_from_files(
    *,
    items_file: str,
    output_path: str,
    result_file: str,
    vendor: str,
    request: TaskRequest,
) -> int:
    started_at = time.monotonic()
    rows: list[dict] = []
    saved_path: str | None = None
    error: str | None = None
    error_type: str | None = None
    error_stage = "runtime_init"
    error_traceback: str | None = None
    items: list[ScrapeItem] = []
    request.output_path = output_path
    request.excel_input = None
    request.input_paths = []
    request.urls_file = None
    request.workers = 1
    if not request.browser_options.session_file:
        request.browser_options.session_file = default_session_path()
    stop_event = FileStopEvent(request.stop_file)

    def reporter(event: dict[str, Any] | str) -> None:
        return None

    scraper = None
    stopped = False
    try:
        raw_items = _read_json(items_file)
        items = [_scrape_item_from_dict(item) for item in raw_items]
        reporter({"event": "process_started", "stage": "process_started", "message": "independent worker process started", "url_count": len(items)})
        if stop_event.is_set():
            stopped = True
        elif not request.browser_options.driver_executable_path:
            error_stage = "driver_prepare"
            _prepare_browser_runtime(request, reporter)
            stopped = stop_event.is_set()
        if not stopped:
            error_stage = "chrome_detect"
            scraper = TripCliScraper(request, reporter, vendor=vendor, worker_id=1, global_total=len(items))
            scraper.prepare_environment()
            error_stage = "browser_start"
            scraper.start_browser()
            error_stage = "scrape"
            rows = scraper.scrape_items(items, stop_event=stop_event, rows=rows)
            stopped = stop_event.is_set()
        if rows:
            error_stage = "export"
            saved_path = export_rows(
                rows,
                output_path,
                rating_only=request.rating_only,
                extract_real_price=request.extract_real_price,
                extract_product_info=request.extract_product_info,
                include_vendor_column=True,
            )
            reporter({"event": "export_done", "rows": len(rows), "output_path": str(Path(saved_path).expanduser().resolve())})
    except Exception as exc:
        error = str(exc)
        error_type = type(exc).__name__
        error_traceback = traceback.format_exc()
        reporter(
            {
                "event": "worker_failed",
                "stage": error_stage,
                "error": error,
                "error_type": error_type,
                "traceback": error_traceback,
                "rows": len(rows),
                "url_count": len(items),
            }
        )
    finally:
        if scraper:
            scraper.close()

    stopped = stopped or stop_event.is_set()
    success = error is None and bool(rows) and not stopped
    result = {
        "event": "vendor_done",
        "vendor": vendor,
        "worker_id": 1,
        "success": success,
        "stopped": stopped,
        "rows": len(rows),
        "output_path": saved_path or output_path,
        "session_file": request.browser_options.session_file,
        "error": "stopped by request" if stopped and error is None else (error if error else (None if rows else "no rows captured")),
        "error_type": error_type,
        "stage": "stopped" if stopped else (error_stage if error else "done"),
        "traceback": error_traceback,
        "url_count": len(items),
        "urls_done": len(items) if rows else 0,
        "elapsed_ms": int((time.monotonic() - started_at) * 1000),
        "worker_dir": str(Path(result_file).expanduser().parent),
        "profile_dir": request.browser_options.user_data_dir,
        "driver_executable_path": request.browser_options.driver_executable_path,
        "stop_file": request.stop_file,
    }
    _write_json(result_file, result)
    return 0 if success else 4


def _worker_request(request: TaskRequest, output_path: str, session_file: str) -> TaskRequest:
    return TaskRequest(
        input_paths=[],
        urls_file=None,
        excel_input=None,
        output_path=output_path,
        output_dir=None,
        log_file=request.log_file,
        stop_file=request.stop_file,
        workers=1,
        startup_stagger=request.startup_stagger,
        stop_grace_seconds=request.stop_grace_seconds,
        rating_only=request.rating_only,
        extract_real_price=request.extract_real_price,
        extract_product_info=request.extract_product_info,
        browser_options=type(request.browser_options)(
            browser=request.browser_options.browser,
            headless=request.browser_options.headless,
            session_file=session_file,
            chrome_path=request.browser_options.chrome_path,
            chrome_version_main=request.browser_options.chrome_version_main,
            driver_executable_path=request.browser_options.driver_executable_path,
            user_data_dir=request.browser_options.user_data_dir,
            timeout=request.browser_options.timeout,
            memory_guard_enabled=request.browser_options.memory_guard_enabled,
            memory_limit_mb=request.browser_options.memory_limit_mb,
            restart_every=request.browser_options.restart_every,
        ),
    )


def _active_summary(states: dict[str, dict[str, Any]]) -> list[str]:
    summary = []
    for vendor, state in states.items():
        current = state.get("current_index") or 0
        total = state.get("total") or 0
        stage = state.get("stage") or "ready"
        summary.append(f"{vendor}:{current}/{total}:{stage}")
    return summary


def _emit_progress(
    reporter: Reporter,
    *,
    states: dict[str, dict[str, Any]],
    completed_groups: int,
    vendor_count: int,
    urls_done: int,
    url_count: int,
    failed: int,
) -> None:
    reporter(
        {
            "event": "progress",
            "vendors_done": completed_groups,
            "vendor_count": vendor_count,
            "urls_done": urls_done,
            "url_count": url_count,
            "active": _active_summary(states),
            "failed": failed,
        }
    )


def _preflight(request: TaskRequest, reporter: Reporter) -> None:
    reporter({"event": "browser_prepare", "message": f"python={sys.executable}"})
    reporter({"event": "browser_prepare", "message": f"platform={platform.platform()}"})
    if request.log_file:
        Path(request.log_file).expanduser().parent.mkdir(parents=True, exist_ok=True)


def _apply_browser_startup_info(request: TaskRequest, info: BrowserStartupInfo) -> None:
    request.browser_options.chrome_path = info.chrome_path
    request.browser_options.chrome_version_main = info.chrome_version_main
    request.browser_options.driver_executable_path = info.driver_executable_path


def _prepare_browser_runtime(request: TaskRequest, reporter: Reporter) -> BrowserStartupInfo:
    reporter({"event": "driver_prepare_started", "stage": "driver_prepare", "message": "preparing chromedriver"})
    try:
        info = prepare_chromedriver()
    except BrowserStartupError as exc:
        event = {
            "event": "driver_prepare_failed" if exc.stage == "driver_prepare" else "startup_preflight_failed",
            "stage": exc.stage,
            "error_type": exc.error_type,
            "error": str(exc),
            "suggestion": exc.suggestion,
            "traceback": exc.traceback_text,
            "log_file": request.log_file,
        }
        event.update(exc.details)
        reporter(event)
        raise
    _apply_browser_startup_info(request, info)
    reporter(
        {
            "event": "driver_prepare_done",
            "stage": "driver_prepare",
            "message": "chromedriver ready",
            "chrome_path": info.chrome_path,
            "chrome_version_main": info.chrome_version_main,
            "driver_executable_path": info.driver_executable_path,
            "driver_reused": info.driver_reused,
            "cache_dir": info.cache_dir,
            "ssl_cert_file": info.ssl_cert_file,
            "ssl_cert_source": info.ssl_cert_source,
            "elapsed_ms": info.driver_prepare_ms,
            "log_file": request.log_file,
        }
    )
    return info


def _scrape_vendor_batches_subprocess(
    request: TaskRequest,
    reporter: Reporter,
    run_mode: str | None = None,
    output_base_dir: str | None = None,
    should_stop: Callable[[], bool] | None = None,
    progress_callback: Callable[[int, int], None] | None = None,
    rows_callback: Callable[[list[dict]], None] | None = None,
) -> TaskResult:
    try:
        items = load_items(request)
        groups = _vendor_groups(items)
        workers = _resolve_worker_count(request.workers, len(groups))
    except Exception as exc:
        return TaskResult(success=False, error=str(exc), exit_code=2)

    warnings = load_excel_items(request.excel_input)[1] if request.excel_input else []
    for warning in warnings:
        reporter({"event": "warning", "message": warning})

    timestamp = datetime.now().strftime("%y%m%d-%H%M")
    if request.output_dir:
        output_dir = _expand_output_path(request.output_dir)
    elif output_base_dir:
        output_dir = str(Path(output_base_dir).expanduser() / "trip_vendor_results" / timestamp)
    else:
        output_dir = default_vendor_output_dir(timestamp=timestamp)
    Path(output_dir).mkdir(parents=True, exist_ok=True)
    if request.log_file:
        Path(request.log_file).expanduser().parent.mkdir(parents=True, exist_ok=True)

    total_urls = sum(len(group["items"]) for group in groups)
    reporter(
        {
            "event": "run_started",
            "mode": "excel",
            "run_mode": run_mode,
            "vendor_count": len(groups),
            "url_count": total_urls,
            "workers": workers,
            "headless": request.browser_options.headless,
            "rating_only": request.rating_only,
            "extract_real_price": request.extract_real_price,
            "extract_product_info": request.extract_product_info,
            "output_dir": str(Path(output_dir).expanduser().resolve()),
            "log_file": str(Path(request.log_file).expanduser().resolve()) if request.log_file else None,
        }
    )
    _preflight(request, reporter)
    try:
        _prepare_browser_runtime(request, reporter)
    except BrowserStartupError as exc:
        failed_vendors = [str(group["vendor"]) for group in groups]
        error = str(exc)
        if exc.suggestion:
            error = f"{error}; suggestion: {exc.suggestion}"
        return TaskResult(False, output_dir, [], 0, warnings, failed_vendors, request.log_file, error, 4)

    base_session = str(Path(request.browser_options.session_file or default_session_path()).expanduser())
    workers_dir = Path(output_dir).expanduser() / "_workers"
    workers_dir.mkdir(parents=True, exist_ok=True)
    pending = list(groups)
    active: list[dict[str, Any]] = []
    active_states: dict[str, dict[str, Any]] = {}
    results: list[dict[str, Any]] = []
    output_paths: list[str] = []
    completed_groups = 0
    urls_done = 0
    stop_requested = {"value": False}
    stop_started_at = {"value": None}
    force_terminate_sent = {"value": False}
    old_handlers = {}
    next_worker_id = {"value": 0}
    stop_attempts = {"value": 0}

    def request_active_workers_stop() -> None:
        for entry in active:
            _write_stop_request(entry.get("stop_file"))

    def terminate_active_workers() -> None:
        for entry in active:
            process = entry["process"]
            _terminate_process_tree(process)

    def stop_all(signum, _frame) -> None:
        stop_attempts["value"] += 1
        if not stop_requested["value"]:
            stop_started_at["value"] = time.monotonic()
        stop_requested["value"] = True
        if stop_attempts["value"] == 1:
            request_active_workers_stop()
            reporter(
                {
                    "event": "stop_requested",
                    "message": f"received signal {signum}; stop files written",
                    "grace_seconds": request.stop_grace_seconds,
                    "mode": "file",
                }
            )
            return
        force_terminate_sent["value"] = True
        reporter({"event": "warning", "message": f"received signal {signum}; terminating worker process trees"})
        terminate_active_workers()

    for sig in (signal.SIGINT, signal.SIGTERM):
        old_handlers[sig] = signal.getsignal(sig)
        signal.signal(sig, stop_all)

    def start_next(group: dict[str, Any]) -> dict[str, Any]:
        next_worker_id["value"] += 1
        worker_id = next_worker_id["value"]
        output_path = request.output_path if len(groups) == 1 and request.output_path else _vendor_output_path(group, output_dir, timestamp)
        spec = _create_worker_spec(
            request=request,
            output_path=output_path,
            group=group,
            worker_id=worker_id,
            workers_dir=workers_dir,
            base_session=base_session,
        )
        reporter({"event": "worker_spawn_started", "vendor": group["vendor"], "worker_id": worker_id, "message": "spawning independent worker"})
        process = subprocess.Popen(spec["command"], **_worker_popen_kwargs())
        reporter(
            {
                "event": "vendor_started",
                "vendor": group["vendor"],
                "worker_id": worker_id,
                "url_count": len(group["items"]),
                "worker_dir": spec["worker_dir"],
            }
        )
        active_states[group["vendor"]] = {"total": len(group["items"]), "current_index": 0, "stage": "running"}
        if request.startup_stagger > 0:
            time.sleep(request.startup_stagger)
        spec["process"] = process
        spec["started_at"] = time.monotonic()
        return spec

    try:
        while pending and len(active) < workers and not stop_requested["value"]:
            active.append(start_next(pending.pop(0)))

        while active:
            if should_stop and should_stop():
                if not stop_requested["value"]:
                    stop_started_at["value"] = time.monotonic()
                    request_active_workers_stop()
                    reporter({"event": "stop_requested", "message": "stop files written", "grace_seconds": request.stop_grace_seconds, "mode": "file"})
                stop_requested["value"] = True
            if stop_requested["value"] and stop_started_at["value"] is not None:
                if (
                    not force_terminate_sent["value"]
                    and request.stop_grace_seconds >= 0
                    and time.monotonic() - stop_started_at["value"] > request.stop_grace_seconds
                ):
                    force_terminate_sent["value"] = True
                    reporter({"event": "warning", "message": "stop grace period exceeded; terminating worker process trees"})
                    terminate_active_workers()

            still_active = []
            for entry in active:
                process = entry["process"]
                if process.poll() is None:
                    still_active.append(entry)
                    continue

                vendor = entry["group"]["vendor"]
                active_states.pop(vendor, None)
                completed_groups += 1
                result_file = Path(entry["result_file"]).expanduser()
                if result_file.exists():
                    result = _read_json(result_file)
                    result["worker_id"] = entry["worker_id"]
                    result.setdefault("vendor", vendor)
                    result.setdefault("output_path", entry["output_path"])
                    result.setdefault("session_file", entry["session_file"])
                    result.setdefault("worker_dir", entry["worker_dir"])
                    result.setdefault("profile_dir", entry["profile_dir"])
                    result.setdefault("driver_executable_path", entry["driver_executable_path"])
                    result.setdefault("stop_file", entry["stop_file"])
                    result.setdefault("url_count", len(entry["group"]["items"]))
                    result.setdefault("exitcode", process.returncode)
                else:
                    result = {
                        "event": "vendor_done",
                        "vendor": vendor,
                        "worker_id": entry["worker_id"],
                        "success": False,
                        "rows": 0,
                        "output_path": entry["output_path"],
                        "session_file": entry["session_file"],
                        "worker_dir": entry["worker_dir"],
                        "profile_dir": entry["profile_dir"],
                        "driver_executable_path": entry["driver_executable_path"],
                        "stop_file": entry["stop_file"],
                        "error": f"worker exited with code {process.returncode} without result file",
                        "error_type": "MissingWorkerResult",
                        "stage": "worker_exit",
                        "url_count": len(entry["group"]["items"]),
                        "urls_done": 0,
                        "exitcode": process.returncode,
                    }
                results.append(result)
                urls_done += int(result.get("url_count") or 0) if result.get("success") else int(result.get("urls_done") or 0)
                if progress_callback:
                    progress_callback(completed_groups, len(groups))
                if result.get("output_path") and (result.get("success") or int(result.get("rows") or 0) > 0):
                    output_paths.append(result["output_path"])
                reporter(result)
                _emit_progress(
                    reporter,
                    states=active_states,
                    completed_groups=completed_groups,
                    vendor_count=len(groups),
                    urls_done=urls_done,
                    url_count=total_urls,
                    failed=len([current for current in results if not current.get("success")]),
                )
                if pending and not stop_requested["value"]:
                    still_active.append(start_next(pending.pop(0)))
            active = still_active
            if active:
                time.sleep(0.5)

        successful_session = next((result.get("session_file") for result in results if result.get("success") and result.get("session_file")), None)
        if successful_session and Path(successful_session).exists():
            base_path = Path(base_session)
            base_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(successful_session, base_path)

        rows = sum(int(result.get("rows") or 0) for result in results)
        failures = [result for result in results if not result.get("success")]
        failed_vendors = [str(result.get("vendor")) for result in failures if result.get("vendor")]
        if stop_requested["value"]:
            return TaskResult(False, output_dir, output_paths, rows, warnings, failed_vendors, request.log_file, "stopped by signal", 130)
        if failures:
            error = "; ".join(f"{failure.get('vendor')}: {failure.get('error')}" for failure in failures)
            return TaskResult(False, output_dir, output_paths, rows, warnings, failed_vendors, request.log_file, error, 4)
        if rows == 0 and groups:
            failed_vendors = [str(group["vendor"]) for group in groups]
            error = "no rows captured and no successful worker output"
            return TaskResult(False, output_dir, output_paths, rows, warnings, failed_vendors, request.log_file, error, 4)
        return TaskResult(True, output_dir, output_paths, rows, warnings, [], request.log_file, None, 0)
    finally:
        for sig, handler in old_handlers.items():
            signal.signal(sig, handler)
        terminate_active_workers()


def scrape_vendor_batches(
    request: TaskRequest,
    reporter: Reporter,
    run_mode: str | None = None,
    output_base_dir: str | None = None,
    should_stop: Callable[[], bool] | None = None,
    progress_callback: Callable[[int, int], None] | None = None,
    rows_callback: Callable[[list[dict]], None] | None = None,
) -> TaskResult:
    return _scrape_vendor_batches_subprocess(
        request,
        reporter,
        run_mode=run_mode,
        output_base_dir=output_base_dir,
        should_stop=should_stop,
        progress_callback=progress_callback,
        rows_callback=rows_callback,
    )


def scrape_xxx(request: TaskRequest, reporter: Reporter, run_mode: str | None = None) -> TaskResult:
    if request.excel_input:
        return scrape_vendor_batches(request, reporter, run_mode=run_mode)

    try:
        items = load_items(request)
    except Exception as exc:
        return TaskResult(success=False, error=str(exc), exit_code=2)

    output_path = _expand_output_path(request.output_path or default_output_path(request))
    if request.log_file:
        Path(request.log_file).expanduser().parent.mkdir(parents=True, exist_ok=True)
    if not request.browser_options.session_file:
        request.browser_options.session_file = default_session_path()
    reporter(
        {
            "event": "run_started",
            "mode": "single",
            "run_mode": run_mode,
            "vendor_count": 1,
            "url_count": len(items),
            "workers": 1,
            "headless": request.browser_options.headless,
            "rating_only": request.rating_only,
            "extract_real_price": request.extract_real_price,
            "extract_product_info": request.extract_product_info,
            "output": str(Path(output_path).expanduser().resolve()),
            "log_file": str(Path(request.log_file).expanduser().resolve()) if request.log_file else None,
        }
    )
    _preflight(request, reporter)
    single_progress = {"done": 0}
    stop_event = multiprocessing.Event()
    old_handlers = {}
    scraper_holder: dict[str, Any] = {}
    stop_attempts = {"value": 0}
    force_stop_requested = {"value": False}

    def stop_single(signum, _frame) -> None:
        stop_attempts["value"] += 1
        if stop_attempts["value"] == 1:
            stop_event.set()
            reporter(
                {
                    "event": "stop_requested",
                    "message": f"received signal {signum}; graceful stop requested",
                    "grace_seconds": request.stop_grace_seconds,
                    "attempt": 1,
                    "mode": "graceful",
                }
            )
            scraper = scraper_holder.get("scraper")
            driver = getattr(scraper, "driver", None)
            if driver:
                try:
                    driver.execute_script("window.stop();")
                except Exception:
                    pass
            return
        force_stop_requested["value"] = True
        reporter(
            {
                "event": "force_stop_requested",
                "message": f"received signal {signum}; forcing browser shutdown",
                "attempt": stop_attempts["value"],
                "mode": "force",
            }
        )
        scraper = scraper_holder.get("scraper")
        if scraper:
            scraper.close()

    for sig in (signal.SIGINT, signal.SIGTERM):
        old_handlers[sig] = signal.getsignal(sig)
        signal.signal(sig, stop_single)

    try:
        _prepare_browser_runtime(request, reporter)
    except BrowserStartupError as exc:
        error = str(exc)
        if exc.suggestion:
            error = f"{error}; suggestion: {exc.suggestion}"
        return TaskResult(False, output_path, [], 0, [], [], request.log_file, error, 4)
    if stop_event.is_set():
        return TaskResult(False, output_path, [], 0, [], [], request.log_file, "stopped by signal", 130)

    def single_reporter(event: dict[str, Any]) -> None:
        if event.get("event") == "url_started":
            event["global_done"] = single_progress["done"]
            event["global_total"] = len(items)
        elif event.get("event") in {"url_done", "url_failed"}:
            single_progress["done"] += 1
            event["global_done"] = single_progress["done"]
            event["global_total"] = len(items)
            reporter(event)
            _emit_progress(
                reporter,
                states={DEFAULT_VENDOR: {"current_index": event.get("url_index"), "total": len(items), "stage": event.get("event").replace("url_", "")}},
                completed_groups=0,
                vendor_count=1,
                urls_done=single_progress["done"],
                url_count=len(items),
                failed=1 if event.get("event") == "url_failed" else 0,
            )
            return
        reporter(event)

    scraper = TripCliScraper(request, single_reporter, vendor=DEFAULT_VENDOR, worker_id=1, global_total=len(items))
    scraper_holder["scraper"] = scraper
    rows: list[dict] = []
    try:
        scraper.prepare_environment()
        scraper.start_browser()
        rows = scraper.scrape_items(items, stop_event=stop_event, rows=rows)
        if rows:
            saved_path = export_rows(
                rows,
                output_path,
                rating_only=request.rating_only,
                extract_real_price=request.extract_real_price,
                extract_product_info=request.extract_product_info,
                include_vendor_column=False,
            )
            event_name = "partial_export_done" if stop_event.is_set() else "export_done"
            reporter({"event": event_name, "rows": len(rows), "output_path": str(Path(saved_path).expanduser().resolve())})
            if stop_event.is_set():
                return TaskResult(False, saved_path, [], len(rows), [], [], request.log_file, "stopped by signal", 130)
            return TaskResult(True, saved_path, [], len(rows), [], [], request.log_file, None, 0)
        if stop_event.is_set():
            return TaskResult(False, output_path, [], 0, [], [], request.log_file, "stopped by signal", 130)
        return TaskResult(False, output_path, [], 0, [], [], request.log_file, "no score/price data captured", 4)
    except Exception as exc:
        reporter({"event": "error", "message": str(exc), "traceback": traceback.format_exc()})
        if rows:
            try:
                saved_path = export_rows(
                    rows,
                    output_path,
                    rating_only=request.rating_only,
                    extract_real_price=request.extract_real_price,
                    extract_product_info=request.extract_product_info,
                    include_vendor_column=False,
                )
                reporter({"event": "partial_export_done", "rows": len(rows), "output_path": str(Path(saved_path).expanduser().resolve())})
                return TaskResult(False, saved_path, [], len(rows), [], [], request.log_file, str(exc), 130 if stop_event.is_set() else 4)
            except Exception:
                pass
        if stop_event.is_set():
            return TaskResult(False, output_path, [], 0, [], [], request.log_file, "stopped by signal", 130)
        return TaskResult(False, output_path, [], 0, [], [], request.log_file, str(exc), 4)
    finally:
        for sig, handler in old_handlers.items():
            signal.signal(sig, handler)
        scraper.close()
