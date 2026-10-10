#!/usr/bin/env python3
import os
from pathlib import Path
import json
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

root=Path(os.environ.get("TRIP_DATA_ROOT", str(Path.home() / "Documents" / "trip-sku-price-scrape"))) / '260831-rating-audit'
out=root/'Trip产品评分_20260831_部分完成.xlsx'
all_rows=json.loads((root/'all_product_scores.json').read_text())
main_rows=json.loads((root/'main_product_stars.json').read_text())
mode2=Path(os.environ.get("TRIP_DATA_ROOT", str(Path.home() / "Documents" / "trip-sku-price-scrape"))) / '260831-mode2-daily/trip_mode2_daily_6site.xlsx'
wbm=load_workbook(mode2,read_only=True,data_only=True);ws=wbm['HK']; it=ws.iter_rows(values_only=True);h=list(next(it)); pid_i=h.index('商品ID'); name_i=h.index('商品名称');url_i=h.index('URL')
name_map={str(r[pid_i]):{'product_name':str(r[name_i] or ''),'url':str(r[url_i] or '')} for r in it if r[pid_i]};wbm.close()
for r in main_rows:
    r.update(name_map.get(str(r['product_id']),{}))

wb=Workbook(); info=wb.active
assert info is not None
info.title='說明'
info.append(['項目','內容'])
for row in [
 ('產生時間','2026-08-31'),
 ('全產品範圍','最新完整 Trip Outerlist 母清單（2026-08-29）：486 個唯一商品；8/30–8/31 Outerlist 任務失敗，未產生完整當日母清單。'),
 ('全產品總評分','486/486 已由 Trip 評論摘要接口即時取得；其中 4 個無公開評分，保持空白。'),
 ('主要產品範圍','六站促銷標籤定時任務（日常抓取目標.xlsx）的 161 個產品。'),
 ('星級計數方法','逐頁讀取 Trip listComments 每筆原始 score 後累加；未由總評分反推。'),
 ('主要產品驗證結果','117/161：五個星級數合計＝Trip 接口評論總數。'),
 ('未完成項目','44/161：Trip HTTP 430 Whaleguard block，星級欄位留空並標記「受風控封鎖」，不以 0 或推算值填補。'),
 ('重新計算平均評分','(5星×5＋4星×4＋3星×3＋2星×2＋1星×1)÷五個星級合計，四捨五入至小數後兩位。'),
]: info.append(row)

allws=wb.create_sheet('所有產品_總評分')
allws.append(['產品ID','產品名稱','目前總評分','評論總數','資料狀態'])
for r in sorted(all_rows,key=lambda x:(x.get('product_name',''),x['product_id'])):
    status='無公開評分' if not r.get('page_score') else '已取得'
    allws.append([str(r['product_id']),r.get('product_name',''),str(r.get('page_score','')),str(r.get('review_total','')),status])

mainws=wb.create_sheet('主要產品_星級明細')
mainws.append(['產品ID','產品名稱','Trip頁面目前總評分','5星數','4星數','3星數','2星數','1星數','評論總數','重新計算平均評分','驗證狀態'])
for r in sorted(main_rows,key=lambda x:(x.get('product_name',''),x['product_id'])):
    if r.get('stars_ok') and not r.get('error'):
        vals=[int(r[f'{n}_star']) for n in ['five','four','three','two','one']]
        den=sum(vals); avg=(Decimal(vals[0]*5+vals[1]*4+vals[2]*3+vals[3]*2+vals[4])/Decimal(den)).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP) if den else None
        mainws.append([str(r['product_id']),r.get('product_name',''),str(r.get('page_score','')),*(str(x) for x in vals),str(r['review_total']),str(avg) if avg is not None else '', '已驗證：星級合計＝評論總數'])
    else:
        mainws.append([str(r['product_id']),r.get('product_name',''),str(r.get('page_score','')),'','','','','',str(r.get('review_total','')),'','受風控封鎖：未填入星級數'])

header_fill=PatternFill('solid',fgColor='1F4E78');warning=PatternFill('solid',fgColor='FFF2CC')
for ws in wb.worksheets:
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]: c.font=Font(color='FFFFFF',bold=True);c.fill=header_fill;c.alignment=Alignment(horizontal='center',vertical='center')
    ws.row_dimensions[1].height=24
    for col in range(1,ws.max_column+1):
        mx=max(len(str(ws.cell(r,col).value or '')) for r in range(1,min(ws.max_row,500)+1))
        ws.column_dimensions[get_column_letter(col)].width=min(max(mx+2,12),55)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.number_format='@';c.alignment=Alignment(vertical='top',wrap_text=True)
if mainws.max_row:
    for row in range(2,mainws.max_row+1):
        if str(mainws.cell(row,11).value).startswith('受風控'):
            for c in mainws[row]: c.fill=warning
wb.save(out)
print(out)
print('all_rows',allws.max_row-1,'main_rows',mainws.max_row-1)
