#!/usr/bin/env python3
"""Recover failed HK daily rows through the headed Go Mode2 CLI, then merge them."""
from __future__ import annotations
import os
import argparse,importlib.util,json,math,subprocess,sys
from datetime import datetime
from pathlib import Path
import openpyxl
ROOT=Path(os.environ.get("TRIP_PROJECT_ROOT", str(Path(__file__).resolve().parents[1])))
COLLECTOR=ROOT/'scripts/trip_outerlist_mode2_api_daily.py'
def load_collector():
 spec=importlib.util.spec_from_file_location('daily_collector',COLLECTOR); m=importlib.util.module_from_spec(spec); assert spec and spec.loader; spec.loader.exec_module(m); return m
def row_maps(ws):
 headers=[str(c.value) for c in ws[2]]; result={}
 for row in ws.iter_rows(min_row=3,values_only=True):
  d={h:'' if i>=len(row) or row[i] is None else str(row[i]) for i,h in enumerate(headers)}; pid=d.get('商品ID','').strip()
  if pid: result[pid]=d
 return headers,result
def mode2_rows(path):
 wb=openpyxl.load_workbook(path,read_only=True,data_only=True); ws=wb['Trip评分数据']; headers=[str(c.value) for c in ws[1]]; out={}
 for row in ws.iter_rows(min_row=2,values_only=True):
  d={h:'' if i>=len(row) or row[i] is None else str(row[i]) for i,h in enumerate(headers)}; pid=d.get('商品ID','').strip()
  if pid: out[pid]=d
 wb.close(); return out
def main():
 p=argparse.ArgumentParser(); p.add_argument('--run-dir',required=True); p.add_argument('--baseline-workbook',required=True); p.add_argument('--workers',type=int,default=5); args=p.parse_args()
 run=Path(args.run_dir); marker=run/'.collection_done'; meta=json.loads(marker.read_text(encoding='utf-8')); output=Path(meta['output_path']); baseline=Path(args.baseline_workbook)
 if not output.is_file() or not baseline.is_file(): raise SystemExit('输出或基线工作簿不存在')
 wb=openpyxl.load_workbook(output,read_only=True,data_only=True); outer_headers,outer=row_maps(wb['Outerlist商品清单']); hk_headers,existing=row_maps(wb['HK']); wb.close()
 failed=[pid for pid,row in existing.items() if row.get('抓取备注','').strip()]
 if not failed: print('无需 Mode2 补跑'); return
 work=run/'mode2-recovery'; work.mkdir(exist_ok=True); binary=work/'trip-score-cdp'
 subprocess.run(['/opt/homebrew/bin/go','build','-o',str(binary),'./cmd/trip-score-cdp'],cwd=ROOT,check=True)
 chunks=[failed[i::args.workers] for i in range(args.workers)]; processes=[]
 for n,chunk in enumerate(chunks,1):
  if not chunk: continue
  d=work/f'w{n}'; d.mkdir(exist_ok=True); urlfile=d/'urls.txt'; urlfile.write_text(''.join(f'https://hk.trip.com/things-to-do/detail/{pid}/\n' for pid in chunk),encoding='utf-8')
  cmd=[str(binary),'-run-mode','mode2','--delay-profile','normal','--urls-file',str(urlfile),'--sites','hk','--output',str(d/'mode2.xlsx'),'--output-dir',str(d),'--log-file',str(d/'terminal_output.txt'),'--log-level','normal','--browser-restart-every','25','--browser-memory-limit-mb','4096']
  processes.append((n,chunk,d,subprocess.Popen(cmd,stdout=(d/'console.log').open('w'),stderr=subprocess.STDOUT,text=True)))
 for n,chunk,d,proc in processes:
  rc=proc.wait()
  if rc or not (d/'mode2.xlsx').is_file(): raise SystemExit(f'Mode2 worker {n} failed (exit={rc})')
 recovered={}
 for n,chunk,d,proc in processes:
  recovered.update(mode2_rows(d/'mode2.xlsx'))
 missing=[pid for pid in failed if pid not in recovered]
 if missing: raise SystemExit('Mode2 未返回失败商品：'+','.join(missing[:20]))
 for pid in failed:
  row=recovered[pid]; old=existing[pid]
  existing[pid].update({'厂商':outer.get(pid,{}).get('厂商') or old.get('厂商') or '未标注','商品ID':pid,'URL':row.get('URL',''),'商品名称':row.get('商品名称',''),'商品资讯':row.get('商品资讯',''),'总促销标签':row.get('总促销标签',''),'月销量':row.get('月销量',''),'年销量':row.get('年销量',''),'总销量':row.get('总销量',''),'评价数':row.get('评价数',''),'总分':row.get('评分',''),'好评率':row.get('好评率',''),'客户服务':row.get('客户服务',''),'方便程度':row.get('方便程度',''),'讯号状态':row.get('讯号状态',''),'中差评价数':row.get('中差评价数',''),'抓取备注':''})
 collector=load_collector(); previous=collector.metrics_from_book(baseline); ordered=[existing[pid] for pid in outer if pid in existing]; collector.add_sales_sheets(output,datetime.now().date(),{'HK':ordered},previous)
 meta.update({'failed_count':0,'failed_samples':[],'baseline_run':str(baseline),'sales_delta_definition':'当次总销量-前一天总销量','recovery_method':'headed_go_mode2_normal','mode2_recovered_count':len(failed),'completed_at':datetime.now().isoformat()}); marker.write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps({'mode2_recovered':len(failed),'output':str(output),'failed_count':0},ensure_ascii=False))
if __name__=='__main__': main()
