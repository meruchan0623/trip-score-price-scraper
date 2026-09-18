#!/usr/bin/env python3
"""One-off browserless Mode2 collector from the public Trip product-info API."""
from __future__ import annotations
import concurrent.futures as futures
import json, re, sys, time, urllib.request
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment

OUTER = Path('/Users/example/Documents/tripcom-outerlist-scraping/260829-0930/items.jsonl')
ATTR = Path('/Users/example/Documents/trip-operations/vendor-attribution/manual_vendor_attribution_latest.json')
ROOT = Path('/Users/example/Documents/trip-sku-price-scrape')
SITES = [('tw','zh-TW','TWD'), ('hk','zh-HK','HKD'), ('ru','ru-RU','RUB'), ('jp','ja-JP','JPY'), ('vn','vi-VN','VND'), ('th','th-TH','THB')]
HEADERS = ['厂商','商品ID','URL','商品名称','商品资讯','总促销标签','月销量','年销量','总销量','评价数','评分','好评率','客户服务','方便程度','讯号状态','中差评价数','抓取备注']
SPACE = re.compile(r'\s+')

def norm(v):
    if v is None: return ''
    return str(v)

def info(add):
    parts=[]
    allowed={'EffectiveDate','FeeInclude','FeeExclude','NetworkType','ApplicableArea'}
    for x in add or []:
        if x.get('code') not in allowed or not x.get('title'): continue
        vals=[SPACE.sub(' ', str(y.get('value','')).strip()) for y in x.get('contents',[]) if y.get('value')]
        if vals: parts.append('• '+str(x['title']).strip()+'：'+' '.join(vals))
    return ' '.join(parts)

def fetch(site, locale, currency, item):
    pid=item['id']; vendor=item['vendor']
    payload={'clientInfo':{'currency':currency,'locale':locale,'pageId':'10650010429','channelId':120,'platformId':26,'oriSyscode':'09','syscode':'09'}, 'enviroment':'PROD', 'searchParams':[{'id':pid,'idType':1}], 'needBasicInfo':True, 'needCommentInfo':True, 'needSalesVolume':True, 'needAddInfos':True, 'needResourceInfos':True, 'needPackageInfos':True, 'needMediaInfo':True, 'needVendorInfo':True, 'needTagInfos':True}
    url=f'https://{site}.trip.com/restapi/soa2/21052/getProductInfo'
    req=urllib.request.Request(url, data=json.dumps(payload,separators=(',',':')).encode(), method='POST', headers={'Content-Type':'application/json','Accept':'application/json','Accept-Language':locale,'User-Agent':'Mozilla/5.0'})
    err=''
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req,timeout=35) as r: data=json.load(r)
            if data.get('ResponseStatus',{}).get('Ack')!='Success': raise RuntimeError((data.get('ResponseStatus',{}).get('Errors') or [{}])[0].get('Message','API failed'))
            p=(data.get('productInfos') or [{}])[0]
            if norm(p.get('productId')) != pid: raise RuntimeError('returned product ID mismatch')
            sale=p.get('salesVolume') or {}; comment=p.get('commentInfo') or {}; basic=p.get('basicInfo') or {}
            return [vendor,pid,f'https://{site}.trip.com/things-to-do/detail/{pid}/?curr={currency}&locale={locale}',norm(basic.get('productName')),info(p.get('addInfos')), '',norm(sale.get('monthlySales')),norm(sale.get('yearlySales')),norm(sale.get('twoYearSales')),norm(comment.get('commentCount')),norm(comment.get('commentScore')),'','','','','','']
        except Exception as e:
            err=f'{type(e).__name__}: {e}'
            if attempt < 2: time.sleep(1.5*(attempt+1))
    return [vendor,pid,f'https://{site}.trip.com/things-to-do/detail/{pid}/?curr={currency}&locale={locale}','','','','','','','','','','','','','',f'接口请求失败：{err}']

def main():
    rows=[json.loads(x) for x in OUTER.open() if x.strip()]
    lookup=json.load(ATTR.open())['lookup_by_product_id']
    seen=set(); items=[]
    for r in rows:
        pid=r['id']
        if pid in seen: continue
        seen.add(pid)
        items.append({'id':pid, 'vendor':lookup.get(pid,{}).get('vendor_name','未标注')})
    run=ROOT/(datetime.now().strftime('%y%m%d-%H%M')+'-mode2-outerlist-api')
    run.mkdir(parents=True,exist_ok=False)
    (run/'任务说明.md').write_text('# Trip Mode2 外部列表六站接口抓取\n\n- 来源：trip-outerlist-daily 260829-0930，共 522 个产品。\n- 站点：TW/HK/RU/JP/VN/TH。\n- 方式：public getProductInfo API，未启动浏览器；并发 5。\n- 厂商：人工标注映射，未命中标为「未标注」。\n',encoding='utf-8')
    workbook=openpyxl.Workbook(); workbook.remove(workbook.active)
    failures=[]; total=len(items)*len(SITES); done=0
    for site,locale,currency in SITES:
        ws=workbook.create_sheet(site.upper()); ws.append(HEADERS)
        for c in ws[1]: c.font=Font(bold=True); c.alignment=Alignment(horizontal='center',vertical='center')
        with futures.ThreadPoolExecutor(max_workers=5) as pool:
            fs=[pool.submit(fetch,site,locale,currency,item) for item in items]
            for f in fs:
                row=f.result(); ws.append(row); done+=1
                if row[-1]: failures.append(f'{site}/{row[1]}')
                if done%25==0 or done==total: print(json.dumps({'event':'progress','done':done,'total':total,'failed':len(failures)},ensure_ascii=False),flush=True)
        ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:Q{ws.max_row}'
        widths=[12,14,50,55,70,40,12,12,12,12,10,12,12,12,12,14,45]
        for idx,width in enumerate(widths,1): ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width=width
    output=run/'trip_mode2_outerlist_6site_api.xlsx'; workbook.save(output)
    state={'run_dir':str(run),'output_path':str(output),'input_products':len(items),'sites':','.join(x[0] for x in SITES),'processed_count':total,'failed_count':len(failures),'failed_samples':failures[:20],'method':'public_product_info_api_no_browser','completed_at':datetime.now().isoformat()}
    (run/'.collection_done').write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'event':'run_done','success':not failures,**state},ensure_ascii=False),flush=True)
    return 0 if not failures else 4
if __name__=='__main__': raise SystemExit(main())
