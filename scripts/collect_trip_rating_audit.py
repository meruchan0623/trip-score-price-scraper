#!/usr/bin/env python3
"""Collect current Trip product scores and exact 1-5 star counts from Trip public APIs."""
from __future__ import annotations

import argparse, csv, json, math, threading, time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Any
import requests

SUMMARY_URL='https://hk.trip.com/restapi/soa2/18084/commonQueryCommentSummary'
LIST_URL='https://hk.trip.com/restapi/soa2/20047/listComments'
HEADERS={'User-Agent':'Mozilla/5.0','Origin':'https://hk.trip.com','Referer':'https://hk.trip.com/','Content-Type':'application/json'}

def req(body: dict[str,Any], url: str, retries=3)->dict[str,Any]:
    last=''
    for i in range(retries):
        try:
            r=requests.post(url,json=body,headers=HEADERS,timeout=35)
            try:
                d=r.json()
            except ValueError:
                last=f'HTTP {r.status_code} non_json: {r.text[:160]!r}'
                time.sleep(1.2*(i+1))
                continue
            if r.status_code==200 and d.get('ret',{}).get('code')==0: return d
            last=f'HTTP {r.status_code}: {d.get("ret",{})}'
        except Exception as e: last=repr(e)
        time.sleep(1.2*(i+1))
    raise RuntimeError(last)

def body(pid:str, page_no:int=1, page_size:int=10):
    return {'ucpBizId':6,'scene':'PRODUCT_QUERY','queryId':int(pid),'paging':{'pageNo':page_no,'pageSize':page_size},'extRequest':{'sourceType':'all'},'tagTerms':[]}

def load_all(path:Path):
    out=[]; seen=set()
    with path.open(encoding='utf-8-sig',newline='') as f:
        for row in csv.DictReader(f):
            pid=str(row.get('id','')).strip()
            if pid and pid not in seen:
                seen.add(pid);out.append({'product_id':pid,'product_name':row.get('name','').strip(),'url':row.get('url','').strip()})
    return out

def load_main(path:Path):
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True);ws=wb['Sheet1']; h=[str(c.value or '').strip() for c in next(ws.iter_rows(max_row=1))]
    vendor=[h.index(x) for x in ['TSIM 網址','BC','TGT','FR','JY','其他'] if x in h]
    dest=h.index('目的地');out=[];seen=set()
    for row in ws.iter_rows(min_row=2,values_only=True):
        for col in vendor:
            v=row[col]
            if v is not None and str(v).strip():
                pid=str(v).strip()
                if pid not in seen:
                    seen.add(pid);out.append({'product_id':pid,'destination':str(row[dest] or '').strip()})
    wb.close();return out

def summary(item):
    d=req(body(item['product_id']),SUMMARY_URL); a=d.get('commentAggregation') or {}
    return {**item,'page_score':str(a.get('scoreAvg','')),'review_total':str(a.get('totalCount','')),'summary_ok':True}

def stars(item):
    # Trip enforces 50 comments per page even when a larger pageSize is sent.
    # Use the server's maximum explicitly so every review is counted once.
    pid=item['product_id']; page_size=50; first=req(body(pid,1,page_size),LIST_URL)
    total=int(first.get('totalCount') or 0); comments=list(first.get('comments') or [])
    pages=math.ceil(total/page_size)
    for page in range(2,pages+1):
        d=req(body(pid,page,page_size),LIST_URL); comments.extend(d.get('comments') or [])
    counts={i:0 for i in range(1,6)}; invalid=[]
    seen=set()
    for c in comments:
        cid=str(c.get('commentId',''))
        if cid in seen: continue
        seen.add(cid)
        try: score=int(float(c.get('score')))
        except: score=0
        if score in counts: counts[score]+=1
        else: invalid.append({'comment_id':cid,'score':c.get('score')})
    agg=(first.get('commentAggregation') or {})
    return {**item,'page_score':str(agg.get('scoreAvg','')),'review_total':str(total),'five_star':str(counts[5]),'four_star':str(counts[4]),'three_star':str(counts[3]),'two_star':str(counts[2]),'one_star':str(counts[1]),'counted_reviews':str(sum(counts.values())),'unclassified_reviews':str(len(invalid)),'pages':str(pages),'stars_ok':sum(counts.values())+len(invalid)==total}

def main():
 p=argparse.ArgumentParser();p.add_argument('--outerlist',required=True);p.add_argument('--targets',required=True);p.add_argument('--out-dir',required=True);p.add_argument('--workers',type=int,default=5);a=p.parse_args()
 out=Path(a.out_dir);out.mkdir(parents=True,exist_ok=True)
 all_items=load_all(Path(a.outerlist));main_items=load_main(Path(a.targets));all_map={x['product_id']:x for x in all_items}
 for x in main_items:
  x.update({k:all_map.get(x['product_id'],{}).get(k,'') for k in ['product_name','url']})
 (out/'input_counts.json').write_text(json.dumps({'all_products':len(all_items),'main_products':len(main_items)},ensure_ascii=False,indent=2))
 def collect(items,fn,path):
  rows=[];lock=threading.Lock()
  with ThreadPoolExecutor(max_workers=a.workers) as ex:
   fs={ex.submit(fn,x):x for x in items}
   for n,f in enumerate(as_completed(fs),1):
    item=fs[f]
    try:r= f.result()
    except Exception as e:r={**item,'error':str(e)}
    with lock:
     rows.append(r); path.write_text(json.dumps(sorted(rows,key=lambda z:z['product_id']),ensure_ascii=False,indent=2))
     print(path.name,n,'/',len(items),item['product_id'],'OK' if 'error' not in r else 'ERR',flush=True)
  return sorted(rows,key=lambda z:z['product_id'])
 collect(all_items,summary,out/'all_product_scores.json')
 collect(main_items,stars,out/'main_product_stars.json')
 print('DONE',out,flush=True)
if __name__=='__main__':main()
