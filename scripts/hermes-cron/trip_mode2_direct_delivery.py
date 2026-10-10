#!/usr/bin/env python3
"""Deliver a completed Mode2 collection once; no agent/model inference involved."""
from __future__ import annotations
import os

import json
import subprocess
import sys
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(os.environ.get("TRIP_DATA_ROOT", str(Path.home() / "Documents" / "trip-sku-price-scrape")))
PROJECT = Path(os.environ.get("TRIP_PROJECT_ROOT", str(Path(__file__).resolve().parents[2])))
DELIVER = PROJECT / "trip_deliver.py"
TARGET = os.environ.get("TRIP_DELIVERY_TARGET", "").strip()
SITES = {"TW", "HK", "RU", "JP", "VN", "TH"}
# The Go CLI defaults 厂商 to this when nothing real was attributed.
LEGACY_VENDOR_PLACEHOLDERS = {"手动输入", "手動輸入"}


def send(text: str) -> None:
    if not TARGET:
        raise RuntimeError("Set TRIP_DELIVERY_TARGET before delivery")
    result = subprocess.run(["hermes", "send", "--to", TARGET, text], capture_output=True, text=True)
    if result.returncode:
        raise RuntimeError(result.stderr.strip() or "hermes send failed")


def split_plaintext(text: str, limit: int = 3500) -> list[str]:
    """Split only at line boundaries so Telegram never truncates a product row."""
    chunks: list[str] = []
    current: list[str] = []
    current_size = 0
    for line in text.splitlines():
        addition = len(line) + (1 if current else 0)
        if current and current_size + addition > limit:
            chunks.append("\n".join(current))
            current = [line]
            current_size = len(line)
        else:
            current.append(line)
            current_size += addition
    if current:
        chunks.append("\n".join(current))
    return chunks


def xlsx_in(run_dir: Path) -> Path | None:
    timestamped = sorted(run_dir.glob("*_trip_mode2_daily_6site.xlsx"))
    if timestamped:
        return timestamped[-1]
    standard = run_dir / "trip_mode2_daily_6site.xlsx"
    if standard.is_file() and standard.stat().st_size:
        return standard
    generic = sorted(run_dir.glob("*mode2*6site*.xlsx"))
    return generic[-1] if generic else None


def tags(path: Path) -> dict[tuple[str, str], tuple[str, str, str]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    found: dict[tuple[str, str], tuple[str, str, str]] = {}
    for sheet in wb.sheetnames:
        site = sheet.upper()
        if site not in SITES:
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        columns = {str(v).strip(): i for i, v in enumerate(header) if v is not None}
        pid_i = columns.get("商品ID")
        promo_i = columns.get("总促销标签")
        name_i = columns.get("商品名称")
        vendor_i = columns.get("厂商")
        if pid_i is None or promo_i is None:
            continue
        for row in rows:
            pid = str(row[pid_i]).strip() if len(row) > pid_i and row[pid_i] else ""
            if not pid:
                continue
            promo = str(row[promo_i]).strip() if len(row) > promo_i and row[promo_i] else ""
            name = str(row[name_i]).strip() if name_i is not None and len(row) > name_i and row[name_i] else ""
            vendor = str(row[vendor_i]).strip() if vendor_i is not None and len(row) > vendor_i and row[vendor_i] else ""
            found[(pid, site)] = (promo, name, vendor)
    wb.close()
    return found


def discount_signature(label: str) -> tuple[int, ...]:
    """Promotion semantics, independent of the locale-specific marketing copy."""
    import re
    # Trip's total-promo text is localized, but percentage discounts are stable.
    # Do not treat title/product-info wording changes as a promotion change.
    # Up to three digits: a 100%-off label must not be silently dropped by a
    # two-digit-only pattern (it would look identical to "no promotion").
    return tuple(sorted({int(value) for value in re.findall(r"(?<![\d.])(\d{1,3})\s*%", label)}))


def discount_text(values: tuple[int, ...]) -> str:
    return " + ".join(f"{value}%" for value in values) if values else "无折扣标签"


def _product_name(pid: str, current: dict, previous: dict) -> str:
    """Prefer a compact HK/TW product name and remove technical suffixes."""
    for source in (current, previous):
        for site in ("HK", "TW", "JP", "TH", "VN", "RU"):
            value = source.get((pid, site), ("", "", ""))
            value = value[1] if len(value) > 1 else ""
            if value:
                import re
                raw_name = value.replace("\\|", "|")
                name = re.split(r"[|｜]", raw_name, maxsplit=1)[0].strip(" \\-")
                return name[:48]
    return "未返回商品名"


def _product_vendor(pid: str, current: dict, previous: dict) -> str:
    """Return a stable vendor label for the product across all six sites."""
    vendors: list[str] = []
    for source in (current, previous):
        for site in ("HK", "TW", "JP", "TH", "VN", "RU"):
            value = source.get((pid, site), ("", "", ""))
            vendor = value[2].strip() if len(value) > 2 else ""
            # Legacy workbooks wrote the CLI placeholder instead of a real vendor;
            # never let it leak into a business report as if it were a vendor name.
            if vendor and vendor not in LEGACY_VENDOR_PLACEHOLDERS and vendor not in vendors:
                vendors.append(vendor)
    if not vendors:
        return "未返回厂商"
    if len(vendors) == 1:
        return vendors[0]
    return "/".join(vendors)


def _promo_label_text(labels: list[str]) -> str:
    """Show source promotion labels without letting one row exhaust a message."""
    values = []
    for label in labels:
        normalized = " ".join(str(label).split())
        if normalized and normalized not in values:
            values.append(normalized)
    text = " / ".join(values) if values else "无"
    return text if len(text) <= 120 else f"{text[:117]}…"


def _report_date(path: Path) -> str:
    try:
        return datetime.strptime(path.parent.name[:6], "%y%m%d").strftime("%Y-%m-%d")
    except ValueError:
        return datetime.now().strftime("%Y-%m-%d")


SITES = ("HK", "TW", "JP", "TH", "VN", "RU")
SITE_ORDER = {site: index for index, site in enumerate(SITES)}
# One Telegram text message is capped at 4096 characters; keep a safety margin
# so a daily report never exceeds two messages and never gets clipped.
MESSAGE_LIMIT = 3750
SINGLE_MESSAGE_LIMIT = MESSAGE_LIMIT  # backwards-compatible alias
MAX_MESSAGES = 2
# Room kept for the "（1/2）" page prefix added when two messages are sent.
PAGE_PREFIX_RESERVE = 8
# Room kept on every page for a worst-case omission tail: without it a page can
# be filled edge to edge and the closing "（未列示 …）" line would have nowhere
# to go, so the reader would never learn that rows were left out.
OMISSION_TAIL_RESERVE = 90
VENDOR_PREFIX = "▪ "
ROW_FORMAT = (
    "行格式：▪ 厂商名 下每行 站点/站点｜旧折扣→新折扣｜产品ID,…"
    "（站点 HK/TW/JP/TH/VN/RU；折扣码 5+50 表示 5% 与 50% 两项折扣，- 表示无折扣；"
    "同厂商/同站点/同折扣变化合并为一行）"
)
XLSX_HINT = "完整明细见同日 XLSX 的「总促销标签」标红行"
SECTIONS = (
    ("▼撤销折扣", "lost", "撤销"),
    ("▼力度变化", "shift", "力度"),
    ("▼新增折扣", "gain", "新增"),
)


def _discount_code(values: tuple[int, ...]) -> str:
    return "+".join(str(value) for value in values) if values else "-"


def _site_text(sites) -> str:
    """Render a site set as slash-separated two-letter codes, e.g. `HK/VN/RU`."""
    return "/".join(sorted((site.upper() for site in sites), key=lambda s: SITE_ORDER.get(s, 99)))


def _pid_sort_key(pid: str):
    return (not pid.isdigit(), int(pid) if pid.isdigit() else pid)


def _change_groups(current, previous):
    """Collapse discount-signature changes into compact groups.

    Returns (changes, rows, groups, changed_pairs): `changes` maps product →
    transitions, `rows` counts (product, transition) pairs per kind, `groups`
    maps kind → group dicts keyed by (vendor name, site set, old→new).
    """
    from collections import defaultdict

    changes: dict[str, dict[tuple[tuple[int, ...], tuple[int, ...]], set[str]]] = defaultdict(
        lambda: defaultdict(set)
    )
    for key in sorted(set(current) | set(previous)):
        pid, site = key
        now_sig = discount_signature(current.get(key, ("", "", ""))[0])
        old_sig = discount_signature(previous.get(key, ("", "", ""))[0])
        if now_sig != old_sig:
            changes[pid][(old_sig, now_sig)].add(site)

    grouped: dict[str, dict[tuple[str, str, str], list[str]]] = defaultdict(lambda: defaultdict(list))
    rows: Counter = Counter()
    changed_pairs: set[tuple[str, str]] = set()
    for pid, transitions in changes.items():
        vendor = _product_vendor(pid, current, previous)
        for (old_sig, now_sig), sites in transitions.items():
            changed_pairs.update((pid, site) for site in sites)
            kind = "lost" if not now_sig else ("gain" if not old_sig else "shift")
            transition = f"{_discount_code(old_sig)}→{_discount_code(now_sig)}"
            grouped[kind][(vendor, _site_text(sites), transition)].append(pid)
            rows[kind] += 1

    groups: dict[str, list[dict]] = {}
    for kind, members in grouped.items():
        items = []
        for (vendor, sites, transition), pids in members.items():
            pids = sorted(pids, key=_pid_sort_key)
            items.append(
                {
                    "kind": kind,
                    "vendor": vendor,
                    "sites": sites,
                    "transition": transition,
                    "pids": pids,
                    "line": f"{sites}｜{transition}｜{','.join(pids)}",
                }
            )
        # Vendors stay together under one heading; inside a vendor the widest
        # reach comes first, then the biggest group, then a stable text order.
        items.sort(
            key=lambda g: (g["vendor"], -g["sites"].count("/"), -len(g["pids"]), g["sites"], g["transition"])
        )
        groups[kind] = items
    return changes, rows, groups, changed_pairs


def _omission_tail(omitted: dict[str, list[dict]]) -> str:
    group_count = sum(len(items) for items in omitted.values())
    pids = {pid for items in omitted.values() for group in items for pid in group["pids"]}
    detail = "；".join(f"{label} {len(items)} 组" for label, items in omitted.items() if items)
    return f"（未列示 {group_count} 组／{len(pids)} 个产品ID：{detail}；{XLSX_HINT}）"


def _transition_messages(current, previous, today, limit: int) -> list[str]:
    """Render at most two compact messages, ordered by business priority.

    Under each section every vendor gets one `▪ <exact vendor name>` heading
    followed by rows `站点/站点｜旧折扣→新折扣｜产品ID,…`; products sharing
    vendor, site set and discount transition are merged into one row so a day
    with hundreds of changes still fits. Sections are emitted in priority
    order (撤销 → 力度变化 → 新增), messages split only at whole-group
    boundaries, and a second message carries its own continuation header and
    row-format line (section and vendor headings are repeated there). If even
    two messages cannot hold every group, the second message closes with the
    exact omitted group / product-ID counts; a third message is never
    produced. Raw localized promotion labels stay in the workbook.
    """
    changes, rows, groups, changed_pairs = _change_groups(current, previous)
    title = f"📊 携程 Mode2 六站促销标签变化｜{_report_date(today)}"
    # Site counts are changed (product, site) pairs, so a product changing on
    # three sites is counted once per site, not once per product.
    site_counts = Counter(site for _pid, site in changed_pairs)
    vendor_counts = Counter(_product_vendor(pid, current, previous) for pid in changes)
    tallies = [
        f"变化 {len(changes)} 产品／{sum(rows.values())} 处：撤销 {rows['lost']}｜力度变化 {rows['shift']}｜新增 {rows['gain']}",
        "站点 " + "｜".join(
            f"{site} {site_counts[site]}" for site in sorted(site_counts, key=lambda s: SITE_ORDER.get(s, 99))
        ),
        "厂商 " + "｜".join(f"{name} {count}" for name, count in sorted(vendor_counts.items())),
    ]
    header = [title, *tallies, ROW_FORMAT, ""]
    if not changes:
        return ["\n".join([title, *tallies]) + "\n无实际促销折扣变化（仅文案或语言变化不计）。\n"]

    # A second message repeats the whole header (title marked 续) plus the same
    # row-format line, so a page read on its own still explains its own rows.
    continuation = [f"{title}（续）", *tallies, ROW_FORMAT, ""]
    budget = limit - PAGE_PREFIX_RESERVE - OMISSION_TAIL_RESERVE
    labels = {kind: label for _title, kind, label in SECTIONS}

    # Each page is a list of (text, group_or_None) so overflow handling can pop
    # whole groups, never partial lines.
    pages: list[list[tuple[str, dict | None]]] = [[(line, None) for line in header]]

    def size(page) -> int:
        # Including the trailing newline every rendered message ends with.
        return sum(len(text) + 1 for text, _group in page)

    def fits(page, *texts) -> bool:
        return size(page) + sum(len(text) + 1 for text in texts) <= budget

    def place(page, texts, group) -> None:
        page.extend((text, group if text is group["line"] else None) for text in texts)

    omitted: dict[str, list[dict]] = {label: [] for _title, _kind, label in SECTIONS}
    overflow = False
    for section_title, kind, label in SECTIONS:
        items = groups.get(kind) or []
        if not items:
            continue
        heading: str | None = f"{section_title} {rows[kind]} 处／{len(items)} 组"
        vendor: str | None = None
        for group in items:
            if overflow:
                # Once a group no longer fits, everything after it in priority
                # order is omitted too, so the shown prefix is always the most
                # actionable part of the day.
                omitted[label].append(group)
                continue
            page = pages[-1]
            vendor_line = f"{VENDOR_PREFIX}{group['vendor']}"
            texts = (
                ([heading] if heading else [])
                + ([vendor_line] if group["vendor"] != vendor else [])
                + [group["line"]]
            )
            if fits(page, *texts):
                place(page, texts, group)
                heading, vendor = None, group["vendor"]
                continue
            if len(pages) < MAX_MESSAGES:
                page = [(line, None) for line in continuation]
                pages.append(page)
                # A new message repeats the section and vendor headings so
                # every row stays readable on its own.
                texts = [heading if heading else f"{section_title}（续）", vendor_line, group["line"]]
                if fits(page, *texts):
                    place(page, texts, group)
                    heading, vendor = None, group["vendor"]
                    continue
            overflow = True
            omitted[label].append(group)

    if any(omitted.values()):
        # The tail must sit on the last message; drop trailing whole groups
        # until it fits, counting each of them as omitted too.
        last = pages[-1]
        while not fits(last, _omission_tail(omitted)):
            _text, group = last.pop()
            if group is not None:
                omitted[labels[group["kind"]]].insert(0, group)
        # Never leave a section or vendor heading dangling above the tail.
        while last and last[-1][1] is None and last[-1][0].startswith(("▼", VENDOR_PREFIX)):
            last.pop()
        last.append((_omission_tail(omitted), None))

    messages = ["\n".join(text for text, _group in page) + "\n" for page in pages]
    if len(messages) > 1:
        messages = [f"（{index}/{len(messages)}）\n{text}" for index, text in enumerate(messages, start=1)]
    return messages


def promo_messages(today: Path, previous: Path | None, limit: int = MESSAGE_LIMIT) -> list[str]:
    """Build the compact promotion-change report as one or two messages."""
    current = tags(today)
    yesterday = tags(previous) if previous else {}
    return _transition_messages(current, yesterday, today, limit)


def main() -> int:
    import os

    if not TARGET:
        raise SystemExit("Set TRIP_DELIVERY_TARGET before delivery")

    # Manual-only override (set by trip_daily_mode2_cron.sh under MODE2_FORCE_RERUN=1):
    # re-deliver the latest XLSX and regenerate the promo report even though the
    # .delivery_done / .promo_change_reported markers already exist.
    force = os.environ.get("MODE2_FORCE_DELIVERY") == "1"
    today = datetime.now().strftime("%y%m%d")
    yesterday = (datetime.now() - timedelta(days=1)).strftime("%y%m%d")
    run_dir = ROOT / f"{today}-mode2-daily"
    done = run_dir / ".collection_done"
    delivery = run_dir / ".delivery_done"
    promo_done = run_dir / ".promo_change_reported"
    if not done.exists():
        raise SystemExit("collection marker missing")
    output = xlsx_in(run_dir)
    if not output:
        raise SystemExit("Mode2 output xlsx missing")
    payload = json.loads(done.read_text())
    if force:
        print(f"[delivery] MODE2_FORCE_DELIVERY=1: re-delivering {output}", file=sys.stderr)
    if force or not delivery.exists():
        # Retry up to 3 times with backoff; mark pending on failure so the next
        # cron tick / manual run can compensate instead of losing the deliverable.
        import time
        pending = run_dir / ".pending_delivery.json"
        last_error = ""
        for attempt in range(1, 4):
            command = [sys.executable, str(DELIVER), str(output), "--to", TARGET, "--subject", "Trip Mode2 Daily"]
            if force:
                # A forced resend is an explicit operator action, so the
                # duplicate-safe guard must be bypassed deliberately.
                command.append("--force")
            result = subprocess.run(command, capture_output=True, text=True)
            if result.returncode == 0:
                delivery.write_text(json.dumps({"delivered_at": datetime.now().isoformat(), "output_path": str(output), "target": TARGET, "processed_count": payload.get("processed_count")}))
                pending.unlink(missing_ok=True)
                last_error = ""
                break
            last_error = (result.stderr or result.stdout or "delivery failed").strip()
            if attempt < 3:
                time.sleep(10 * attempt)
        if last_error:
            pending.write_text(json.dumps({"status": "pending", "path": str(output), "target": TARGET, "error": last_error, "attempts": 3}))
            raise SystemExit(last_error)
    if force or not promo_done.exists():
        baseline = xlsx_in(ROOT / f"{yesterday}-mode2-daily")
        if baseline:
            # The renderer already bounds the report to at most two whole-group
            # pages (prefixed "（1/2）"/"（2/2）"); never re-split it with
            # split_plaintext.
            messages = promo_messages(output, baseline)
            for text in messages:
                send(text)
            promo_done.write_text(json.dumps({
                "reported_at": datetime.now().isoformat(),
                "target": TARGET,
                "messages": len(messages),
                "chunks": len(messages),
            }))
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
