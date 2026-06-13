#!/usr/bin/env python3
"""Daily launcher for Trip SKU vendor price trend scrape.

Starts the detached Go CDP scrape quickly. A separate watcher/final-delivery
path handles progress summaries and report delivery.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

SYSTEM_PYTHON = '/usr/local/bin/python3'
if Path(SYSTEM_PYTHON).exists() and Path(sys.executable).resolve() != Path(SYSTEM_PYTHON).resolve():
    os.execv(SYSTEM_PYTHON, [SYSTEM_PYTHON, __file__, *sys.argv[1:]])

from openpyxl import load_workbook

ROOT = Path('/Users/example/Documents/trip-sku-price-scrape')
CONFIG_INPUT = ROOT / 'config' / '非TSIM厂商价格抓取输入.xlsx'
REPO = Path('/Users/example/Develop/trip-score-price-scraper')
BINARY = Path('/tmp/trip-score-cdp')
FINALIZER = REPO / 'scripts' / 'hermes-cron' / 'trip_sku_trend_finalize.py'
FINAL_DELIVERY = Path('/Users/example/.hermes/scripts/trip_sku_non_tsim_final_delivery.py')
DEFAULT_COMPARE_VENDOR = 'TSIM 網址'

POLL_INTERVAL = 1200  # 20 minutes
DONE_RE = re.compile(r'^\[(?P<v>[^\]]+)\] url done (?P<i>\d+)/(?P<t>\d+)')
RISK_RE = re.compile(r'risk challenge detected')


def count_vendors(path: Path):
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb.active
    result = {}
    for col in range(1, (ws.max_column or 1) + 1):
        vendor = ws.cell(1, col).value
        if not vendor:
            continue
        vendor = str(vendor).strip()
        n = 0
        for row in range(2, (ws.max_row or 1) + 1):
            val = ws.cell(row, col).value
            if val and 'trip.com' in str(val):
                n += 1
        if n:
            result[vendor] = n
    return result


def read_state(run_dir: Path):
    p = run_dir / 'current_state.json'
    if not p.exists():
        return {}
    try:
        return json.loads(p.read_text(encoding='utf-8'))
    except Exception:
        return {}


def parse_progress(run_dir: Path, vendor_totals: dict):
    log = run_dir / 'terminal_output.txt'
    done = {}
    risk_count = 0
    if not log.exists():
        return done, risk_count
    with log.open('r', encoding='utf-8', errors='ignore') as f:
        for line in f:
            m = DONE_RE.match(line.rstrip('\n'))
            if m:
                done[m.group('v')] = (int(m.group('i')), int(m.group('t')))
            if RISK_RE.search(line):
                risk_count += 1
    return done, risk_count


def format_progress(run_id: str, elapsed_mins: int, done: dict, vendor_totals: dict, total_urls: int, risk_count: int):
    total_done = sum(v[0] for v in done.values())
    lines = [
        f'⏳ SKU价格抓取进行中 [{run_id}]',
        f'已运行: {elapsed_mins}min | URL: {total_done}/{total_urls} | 反爬拦截: {risk_count}次',
    ]
    for vendor, total in vendor_totals.items():
        if vendor in done:
            cur, _ = done[vendor]
            pct = int(cur / total * 100) if total else 0
            lines.append(f'  {vendor}: {cur}/{total} ({pct}%)')
        else:
            lines.append(f'  {vendor}: 0/{total} (等待中)')
    return '\n'.join(lines)


def format_final(run_dir: Path, run_id: str, state: dict, done: dict, vendor_totals: dict, total_urls: int, risk_count: int):
    exit_path = run_dir / 'exit_code.txt'
    code = exit_path.read_text(errors='ignore').strip() if exit_path.exists() else '?'
    total_done = sum(v[0] for v in done.values())

    start_t = state.get('start_time', 0)
    end_t = state.get('end_time', time.time())
    mins = int((end_t - start_t) / 60) if start_t else 0

    summary_path = run_dir / '趋势对比摘要.json'
    trend = {}
    if summary_path.exists():
        try:
            trend = json.loads(summary_path.read_text(encoding='utf-8'))
        except Exception:
            pass

    lines = [
        f'📊 SKU价格抓取完成 [{run_id}]',
        f'退出码: {code} | 耗时: {mins}min | URL: {total_done}/{total_urls} | 反爬拦截: {risk_count}次',
        f'基准: {trend.get("previous_run_dir") or "无"}',
    ]
    for vendor in sorted(done.keys()):
        cur, tot = done[vendor]
        lines.append(f'  {vendor}: {cur}/{tot}')

    if trend.get('missing_sku_rows') is not None:
        lines.append(f'缺失SKU: {trend["missing_sku_rows"]} | 新增: {trend.get("new_rows", 0)}')

    # Find ZIP package
    pkgs = sorted(run_dir.glob('Trip_SKU厂商价格趋势_*.zip'))
    if not pkgs:
        pkgs = sorted(run_dir.glob('Trip_SKU非TSIM厂商价格趋势_*.zip'))
    if pkgs:
        lines.append('')
        lines.append(f'MEDIA:{pkgs[-1]}')

    return '\n'.join(lines)


def main():
    if not CONFIG_INPUT.exists():
        raise SystemExit(f'配置输入文件不存在：{CONFIG_INPUT}')
    ROOT.mkdir(parents=True, exist_ok=True)
    run_id = time.strftime('%y%m%d-%H%M')
    run_dir = ROOT / run_id
    if run_dir.exists():
        i = 1
        while (ROOT / f'{run_id}-{i}').exists():
            i += 1
        run_dir = ROOT / f'{run_id}-{i}'
    run_dir.mkdir(parents=True, exist_ok=False)
    input_path = run_dir / 'SKU厂商价格趋势输入.xlsx'
    shutil.copy2(CONFIG_INPUT, input_path)
    vendors = count_vendors(input_path)
    total_urls = sum(vendors.values())
    summary = {
        'run_id': run_dir.name,
        'run_dir': str(run_dir),
        'input': str(input_path),
        'config_input': str(CONFIG_INPUT),
        'vendors': vendors,
        'total_urls': total_urls,
        'config': {'run_mode': 'mode1', 'headless': False, 'delay_profile': 'normal',
                    'workers': min(5, len(vendors)), 'log_level': 'verbose',
                    'default_compare_vendor': DEFAULT_COMPARE_VENDOR},
        'status': 'preparing',
        'created_at': time.strftime('%Y-%m-%d %H:%M:%S %z'),
    }
    (run_dir / 'input_summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    (run_dir / '任务说明.md').write_text(
        '# Trip SKU Price Scrape Run — SKU 厂商每日价格趋势抓取\n\n'
        f'- 任务：每天 UTC+8 08:00 抓取配置表中的 Trip.com SKU/外显价/实价/评分/预定数，并与前一次同厂商、同产品、同名 SKU 结果做趋势对比。\n'
        f'- 输入：`{input_path.name}`，来源 `{CONFIG_INPUT}`。\n'
        f'- 厂商范围：' + '，'.join(f'{v}({n})' for v, n in vendors.items()) + f'，合计 {total_urls} 条 URL。\n'
        f'- 参数：`mode1`，有头浏览器，`normal` 延迟，`workers={summary["config"]["workers"]}`。\n'
        f'- 默认日常比价标的：`{DEFAULT_COMPARE_VENDOR}`；这是默认基准厂商，不代表该 cron 只服务某一家厂商，后续可通过配置更换。\n'
        f'- 输出：每厂商原始 XLSX、每厂商含趋势 XLSX、SKU 厂商趋势总表、增强版 vendor_compare 横向比价表（总表为SKU粒度，单链接均价对比为目的地粒度）、趋势摘要 JSON、ZIP 包。\n'
        f'- 说明：预定数变化为两次抓取之间 Trip 页面累计预定数差值，可作为销量动态参考；SKU 比价按结构化规格解析匹配（有效期+总量/每日+流量MB），兼容總量-10GB/總量 10GB/5G-總量-10GB 等格式差异；比价差值只基于匹配SKU交集。\n',
        encoding='utf-8'
    )

    # Build CLI
    subprocess.run(['go', 'build', '-o', str(BINARY), './cmd/trip-score-cdp'], cwd=str(REPO), check=True)
    # Do not persist trip-score-cdp --help output in scheduled runs; it is
    # stable CLI metadata and the user requested not to retain it.

    log_path = run_dir / 'terminal_output.txt'
    runner = run_dir / 'run_daily_sku_vendor_trend.sh'
    workers = summary['config']['workers']
    runner.write_text(f'''#!/bin/bash
set -euo pipefail
RUN_DIR={str(run_dir)!r}
INPUT={str(input_path)!r}
LOG={str(log_path)!r}
echo $$ > "$RUN_DIR/runner.pid"
/usr/bin/env python3 - <<'PY'
import json,time
from pathlib import Path
run=Path({str(run_dir)!r})
state=json.loads((run/'input_summary.json').read_text())
state.update({{'status':'running','start_time':time.time(),'pid':int((run/'runner.pid').read_text().strip()),'log_path':{str(log_path)!r}}})
(run/'current_state.json').write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
Path('/Users/example/Documents/trip-sku-price-scrape/current_state.json').write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
PY
set +e
{str(BINARY)!r} --run-mode mode1 --no-headless --delay-profile normal --excel-input "$INPUT" --workers {workers} --log-level verbose --compare-vendor {DEFAULT_COMPARE_VENDOR!r} --output-dir "$RUN_DIR" > "$LOG" 2>&1
code=$?
set -e
echo "$code" > "$RUN_DIR/exit_code.txt"
echo "DONE exit_code=$code $(date)" >> "$LOG"
/usr/bin/env python3 - <<'PY'
import json,time
from pathlib import Path
run=Path({str(run_dir)!r})
state=json.loads((run/'current_state.json').read_text())
code=int((run/'exit_code.txt').read_text().strip())
state.update({{'status':'scraped' if code == 0 else 'failed','exit_code':code,'scrape_end_time':time.time()}})
(run/'current_state.json').write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
Path('/Users/example/Documents/trip-sku-price-scrape/current_state.json').write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
PY
if [ "$code" -ne 0 ]; then
  exit "$code"
fi
/usr/bin/env python3 {str(FINALIZER)!r} "$RUN_DIR" > "$RUN_DIR/finalize_output.json" 2> "$RUN_DIR/finalize_error.log" || true
/usr/bin/env python3 - <<'PY'
import json,time
from pathlib import Path
run=Path({str(run_dir)!r})
state=json.loads((run/'current_state.json').read_text())
state.update({{'status':'finalized','end_time':time.time()}})
(run/'current_state.json').write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
Path('/Users/example/Documents/trip-sku-price-scrape/current_state.json').write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
PY
/usr/bin/env python3 {str(FINAL_DELIVERY)!r} "$RUN_DIR" > "$RUN_DIR/final_delivery_output.json" 2> "$RUN_DIR/final_delivery_error.log" || true
exit $code
''', encoding='utf-8')
    runner.chmod(0o755)

    # Launch scrape
    proc = subprocess.Popen([str(runner)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
    summary.update({'status': 'launched', 'pid': proc.pid, 'runner': str(runner), 'log_path': str(log_path)})
    (run_dir / 'current_state.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    (ROOT / 'current_state.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')

    print(f'SKU价格抓取已启动 [{run_id}] PID:{proc.pid} | ' +
          ' '.join(f'{v}:{n}' for v, n in vendors.items()) + f' 合计{total_urls}条')

    # Important: Hermes cron has a short per-run execution budget. The launcher must
    # start the detached scrape and exit quickly; progress/final delivery is handled
    # by the daily watcher on its own recurring cron.
    return


if __name__ == '__main__':
    main()
