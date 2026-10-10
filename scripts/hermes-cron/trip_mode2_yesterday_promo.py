#!/usr/bin/env python3
"""Insert a `昨日总促销标签` column before `总促销标签` in a Mode2 six-site workbook.

Usage:
    trip_mode2_yesterday_promo.py CURRENT_XLSX [--root RUN_DIR_BASE]

The baseline is the previous calendar day's finalized workbook (derived from
the run-dir name, e.g. 260911-mode2-daily -> 260910-mode2-daily). The exact
previous-day `总促销标签` string is copied by (sheet, 商品ID); anything missing
yields an empty cell. Re-running never adds a second baseline column.
"""
from __future__ import annotations
import os

import argparse
import sys
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(os.environ.get("TRIP_DATA_ROOT", str(Path.home() / "Documents" / "trip-sku-price-scrape")))
SITES = ("TW", "HK", "RU", "JP", "VN", "TH")
PID_HEADER = "商品ID"
PROMO_HEADER = "总促销标签"
BASELINE_HEADER = "昨日总促销标签"
CHANGED_FONT_COLOR = "FFFF0000"  # red font on today's 总促销标签 when it differs from yesterday


def xlsx_in(run_dir: Path) -> Path | None:
    """Same lookup order as trip_mode2_direct_delivery.xlsx_in (timestamped, legacy, generic)."""
    if not run_dir.is_dir():
        return None
    timestamped = sorted(run_dir.glob("*_trip_mode2_daily_6site.xlsx"))
    if timestamped:
        return timestamped[-1]
    standard = run_dir / "trip_mode2_daily_6site.xlsx"
    if standard.is_file() and standard.stat().st_size:
        return standard
    generic = sorted(run_dir.glob("*mode2*6site*.xlsx"))
    return generic[-1] if generic else None


def previous_run_dir(current_xlsx: Path, root: Path | None = None) -> Path | None:
    """Return the previous calendar day's run dir for a workbook in YYMMDD-mode2-daily/."""
    run_dir = current_xlsx.parent
    try:
        day = datetime.strptime(run_dir.name[:6], "%y%m%d")
    except ValueError:
        return None
    base = root if root is not None else run_dir.parent
    return base / f"{(day - timedelta(days=1)).strftime('%y%m%d')}-mode2-daily"


def _header_index(header: list) -> dict[str, int]:
    return {str(v).strip(): i for i, v in enumerate(header) if v is not None}


def _cell_text(row: tuple, index: int | None) -> str:
    if index is None or index >= len(row) or row[index] is None:
        return ""
    return str(row[index])


def baseline_labels(path: Path) -> dict[tuple[str, str], str]:
    """Map (SITE, 商品ID) -> exact 总促销标签 text from a finalized workbook."""
    import openpyxl

    found: dict[tuple[str, str], str] = {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in wb.sheetnames:
            site = sheet.upper()
            if site not in SITES:
                continue
            rows = wb[sheet].iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            columns = _header_index(list(header))
            pid_i, promo_i = columns.get(PID_HEADER), columns.get(PROMO_HEADER)
            if pid_i is None or promo_i is None:
                continue
            for row in rows:
                pid = _cell_text(row, pid_i).strip()
                if pid:
                    found[(site, pid)] = _cell_text(row, promo_i)
    finally:
        wb.close()
    return found


def enrich(current_xlsx: Path, baseline_xlsx: Path | None) -> dict[str, int]:
    """Insert/refresh the baseline column in place. Returns per-sheet matched counts."""
    import copy

    import openpyxl
    from openpyxl.styles import Font

    labels = baseline_labels(baseline_xlsx) if baseline_xlsx and baseline_xlsx.is_file() else {}
    wb = openpyxl.load_workbook(current_xlsx)
    stats: dict[str, int] = {}
    for sheet in wb.sheetnames:
        site = sheet.upper()
        if site not in SITES:
            continue
        ws = wb[sheet]
        header = [c.value for c in ws[1]] if ws.max_row >= 1 else []
        columns = _header_index(header)
        pid_i, promo_i = columns.get(PID_HEADER), columns.get(PROMO_HEADER)
        if pid_i is None or promo_i is None:
            continue
        base_i = columns.get(BASELINE_HEADER)
        if base_i is None:
            # openpyxl columns are 1-based; insert directly before 总促销标签.
            ws.insert_cols(promo_i + 1)
            base_i = promo_i
            ws.cell(row=1, column=base_i + 1, value=BASELINE_HEADER)
            if pid_i > promo_i:
                pid_i += 1
        matched = 0
        for r in range(2, ws.max_row + 1):
            pid_value = ws.cell(row=r, column=pid_i + 1).value
            pid = str(pid_value).strip() if pid_value is not None else ""
            label = labels.get((site, pid), "") if pid else ""
            ws.cell(row=r, column=base_i + 1, value=label if label else None)
            if label:
                matched += 1
            # Red font on today's 总促销标签 only when a non-empty baseline differs from it;
            # otherwise clear any red left by an earlier run. Only the colour is touched.
            promo_cell = ws.cell(row=r, column=base_i + 2)
            current = str(promo_cell.value) if promo_cell.value is not None else ""
            changed = bool(label) and current != label
            want = CHANGED_FONT_COLOR if changed else None
            current_color = promo_cell.font.color.rgb if promo_cell.font and promo_cell.font.color else None
            if changed != (current_color == CHANGED_FONT_COLOR):
                font = copy.copy(promo_cell.font) if promo_cell.font else Font()  # unwraps the StyleProxy
                font.color = want
                promo_cell.font = font
        stats[site] = matched
    wb.save(current_xlsx)
    return stats


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("current_xlsx", type=Path)
    parser.add_argument("--root", type=Path, default=None, help="run-dir base (default: parent of the run dir)")
    parser.add_argument("--baseline", type=Path, default=None, help="explicit baseline workbook (overrides lookup)")
    args = parser.parse_args(argv)

    current = args.current_xlsx
    if not current.is_file():
        print(f"[yesterday-promo] current workbook missing: {current}", file=sys.stderr)
        return 1
    baseline = args.baseline
    if baseline is None:
        prev_dir = previous_run_dir(current, args.root)
        baseline = xlsx_in(prev_dir) if prev_dir else None
    if baseline is None:
        print("[yesterday-promo] no previous-day workbook found; baseline column left empty", file=sys.stderr)
    stats = enrich(current, baseline)
    print(f"[yesterday-promo] baseline={baseline} matched={stats}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
