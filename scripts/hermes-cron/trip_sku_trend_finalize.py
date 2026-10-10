#!/usr/bin/env python3
"""Append previous-run trend columns to Trip SKU vendor XLSX outputs and package them."""
from __future__ import annotations

import json
import os
import re
import sys
import time
import zipfile
import unicodedata
from collections import defaultdict, deque
from pathlib import Path
from typing import Any

SYSTEM_PYTHON = '/usr/local/bin/python3'
if Path(SYSTEM_PYTHON).exists() and Path(sys.executable).resolve() != Path(SYSTEM_PYTHON).resolve():
    # Hermes cron may invoke scripts with a venv Python that lacks openpyxl.
    os.execv(SYSTEM_PYTHON, [SYSTEM_PYTHON, __file__, *sys.argv[1:]])

from openpyxl import Workbook, load_workbook

ROOT = Path(os.environ.get("TRIP_DATA_ROOT", str(Path.home() / "Documents" / "trip-sku-price-scrape")))
DEFAULT_VENDORS = ['TSIM 網址', 'BC', 'TGT', 'FR', 'JY', '其他']
TREND_COLUMNS = [
    '对比基准日期', '对比基准目录', '趋势匹配状态',
    '前次外显价格', '外显价格变化', '外显价格变化百分比', '外显价格趋势',
    '前次实价', '实价变化', '实价变化百分比', '实价趋势',
    '前次促销标签', '促销标签变化',
    '前次预定数', '预定数变化', '预定数变化百分比', '预定数趋势',
    '前次评价数', '评价数变化', '评价数变化百分比',
    '前次评分', '评分变化', '评分趋势',
    '前次客户服务', '客户服务变化',
    '前次方便程度', '方便程度变化',
    '前次讯号状态', '讯号状态变化',
    '前次中差评价数', '中差评价数变化',
]


def norm_text(value: Any) -> str:
    if value is None:
        return ''
    s = unicodedata.normalize('NFKC', str(value)).strip().lower()
    s = s.replace('－', '-').replace('–', '-').replace('—', '-')
    s = re.sub(r'\s+', '', s)
    return s


def product_id_from_url(url: Any) -> str:
    if not url:
        return ''
    m = re.search(r'/detail/(\d+)', str(url))
    return m.group(1) if m else ''


def num(value: Any):
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = unicodedata.normalize('NFKC', str(value))
    s = s.replace(',', '')
    m = re.search(r'-?\d+(?:\.\d+)?', s)
    return float(m.group(0)) if m else None


def fmt_delta(cur, prev):
    c, p = num(cur), num(prev)
    if c is None or p is None:
        return '', '', '无法比较'
    d = c - p
    pct = '' if p == 0 else d / p
    if d > 0:
        trend = '上涨/增加'
    elif d < 0:
        trend = '下降/减少'
    else:
        trend = '持平'
    return d, pct, trend


def read_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value or '' for c in range(1, (ws.max_column or 1) + 1)]
    rows: list[dict[str, Any]] = []
    carry: dict[str, Any] = {}
    sku_cols = [h for h in headers if str(h).startswith('SKU层级')]
    carry_cols = ['厂商', '商品ID', 'URL', '商品名称', '预定数', '评价数', '评分', '客户服务', '方便程度', '讯号状态', '中差评价数', '商品资讯']
    for r in range(2, (ws.max_row or 1) + 1):
        vals = {headers[c - 1]: ws.cell(r, c).value for c in range(1, len(headers) + 1)}
        if not any(v is not None and str(v) != '' for v in vals.values()):
            continue
        for c in carry_cols:
            if vals.get(c) not in (None, ''):
                carry[c] = vals.get(c)
            elif c in carry:
                vals[c] = carry[c]
        if not vals.get('商品ID'):
            vals['商品ID'] = product_id_from_url(vals.get('URL'))
        vals['_sku_key'] = '|'.join(norm_text(vals.get(c)) for c in sku_cols if norm_text(vals.get(c))) or '无SKU选项'
        vals['_product_key'] = str(vals.get('商品ID') or product_id_from_url(vals.get('URL')) or norm_text(vals.get('URL')))
        vals['_vendor'] = str(vals.get('厂商') or '').strip()
        vals['_values'] = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
        rows.append(vals)
    return [str(h) for h in headers], rows


def vendor_file(run_dir: Path, vendor: str, trend: bool = False) -> Path | None:
    suffix = '_含趋势' if trend else ''
    files = sorted(run_dir.glob(f'{vendor}_*{suffix}.xlsx'))
    files = [p for p in files if '趋势总表' not in p.name]
    if not trend:
        files = [p for p in files if '_含趋势' not in p.stem]
    return files[-1] if files else None


def vendors_for_run(run_dir: Path) -> list[str]:
    """Return configured vendors for a run, preserving config order when available."""
    for name in ('current_state.json', 'input_summary.json'):
        p = run_dir / name
        if not p.exists():
            continue
        try:
            state = json.loads(p.read_text(encoding='utf-8'))
        except Exception:
            continue
        vendors = state.get('vendors') or {}
        if isinstance(vendors, dict) and vendors:
            return [str(v).strip() for v in vendors.keys() if str(v).strip()]
    discovered = [v for v in DEFAULT_VENDORS if vendor_file(run_dir, v)]
    return discovered or DEFAULT_VENDORS


def is_success_run(d: Path) -> bool:
    if not d.is_dir():
        return False
    if not (d / 'exit_code.txt').exists():
        return False
    if (d / 'exit_code.txt').read_text(errors='ignore').strip() != '0':
        return False
    return any(vendor_file(d, v) for v in vendors_for_run(d))


def find_previous_run(current: Path) -> Path | None:
    candidates = []
    for d in ROOT.iterdir() if ROOT.exists() else []:
        if d == current or not d.is_dir() or d.name == 'config':
            continue
        if d.name >= current.name:
            continue
        if is_success_run(d):
            candidates.append(d)
    return sorted(candidates, key=lambda p: p.name)[-1] if candidates else None


def build_prev_index(prev_dir: Path | None):
    idx = defaultdict(deque)
    product_keys = set()
    product_meta = {}
    if not prev_dir:
        return idx, product_keys, product_meta
    for vendor in vendors_for_run(prev_dir):
        p = vendor_file(prev_dir, vendor, trend=False)
        if not p:
            continue
        _, rows = read_rows(p)
        for row in rows:
            key = (vendor, row['_product_key'], row['_sku_key'])
            idx[key].append(row)
            product_keys.add((vendor, row['_product_key']))
            product_meta[(vendor, row['_product_key'])] = row
    return idx, product_keys, product_meta


def trend_for_row(vendor: str, row: dict[str, Any], prev_dir: Path | None, prev_index, prev_product_keys) -> tuple[list[Any], str]:
    base_date = prev_dir.name if prev_dir else ''
    base_dir = str(prev_dir) if prev_dir else ''
    if not prev_dir:
        return [base_date, base_dir, '无前次基准'] + [''] * (len(TREND_COLUMNS) - 3), '无前次基准'
    key = (vendor, row['_product_key'], row['_sku_key'])
    prev = prev_index[key].popleft() if prev_index.get(key) else None
    if prev:
        status = '匹配成功'
    elif (vendor, row['_product_key']) in prev_product_keys:
        status = '新SKU'
    else:
        status = '新产品'
    if not prev:
        return [base_date, base_dir, status] + [''] * (len(TREND_COLUMNS) - 3), status
    display_delta, display_pct, display_trend = fmt_delta(row.get('外显价格'), prev.get('外显价格'))
    real_delta, real_pct, real_trend = fmt_delta(row.get('实价'), prev.get('实价'))
    booking_delta, booking_pct, booking_trend = fmt_delta(row.get('预定数'), prev.get('预定数'))
    review_delta, review_pct, _ = fmt_delta(row.get('评价数'), prev.get('评价数'))
    rating_delta, _, rating_trend = fmt_delta(row.get('评分'), prev.get('评分'))
    cs_delta, _, _ = fmt_delta(row.get('客户服务'), prev.get('客户服务'))
    convenience_delta, _, _ = fmt_delta(row.get('方便程度'), prev.get('方便程度'))
    signal_delta, _, _ = fmt_delta(row.get('讯号状态'), prev.get('讯号状态'))
    bad_delta, _, _ = fmt_delta(row.get('中差评价数'), prev.get('中差评价数'))
    promo_change = ''
    if str(row.get('促销标签') or '') != str(prev.get('促销标签') or ''):
        promo_change = f"{prev.get('促销标签') or ''} -> {row.get('促销标签') or ''}"
    vals = [
        base_date, base_dir, status,
        prev.get('外显价格'), display_delta, display_pct, display_trend,
        prev.get('实价'), real_delta, real_pct, real_trend,
        prev.get('促销标签'), promo_change,
        prev.get('预定数'), booking_delta, booking_pct, booking_trend,
        prev.get('评价数'), review_delta, review_pct,
        prev.get('评分'), rating_delta, rating_trend,
        prev.get('客户服务'), cs_delta,
        prev.get('方便程度'), convenience_delta,
        prev.get('讯号状态'), signal_delta,
        prev.get('中差评价数'), bad_delta,
    ]
    return vals, status


def write_sheet(path: Path, headers: list[str], rows: list[list[Any]], title='趋势明细'):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


def main(run_dir_arg: str):
    run_dir = Path(run_dir_arg)
    if not run_dir.exists():
        raise SystemExit(f'run dir not found: {run_dir}')
    prev_dir = find_previous_run(run_dir)
    prev_index, prev_product_keys, _ = build_prev_index(prev_dir)
    all_headers = None
    all_rows: list[list[Any]] = []
    missing_rows: list[list[Any]] = []
    new_rows: list[list[Any]] = []
    summary = {
        'run_dir': str(run_dir),
        'previous_run_dir': str(prev_dir) if prev_dir else None,
        'generated_at': time.strftime('%Y-%m-%d %H:%M:%S %z'),
        'vendors': {},
    }
    seen_current = set()
    for vendor in vendors_for_run(run_dir):
        src = vendor_file(run_dir, vendor, trend=False)
        if not src:
            continue
        headers, rows = read_rows(src)
        out_headers = headers + TREND_COLUMNS
        out_rows = []
        counts = defaultdict(int)
        for row in rows:
            trend_vals, status = trend_for_row(vendor, row, prev_dir, prev_index, prev_product_keys)
            final = row['_values'] + trend_vals
            out_rows.append(final)
            all_rows.append(final)
            seen_current.add((vendor, row['_product_key'], row['_sku_key']))
            counts[status] += 1
            if status in ('新产品', '新SKU'):
                new_rows.append(final)
        trend_path = run_dir / f'{vendor}_{run_dir.name}_含趋势.xlsx'
        write_sheet(trend_path, out_headers, out_rows, title=f'{vendor}趋势')
        all_headers = out_headers
        summary['vendors'][vendor] = {'source': str(src), 'trend_file': str(trend_path), 'rows': len(out_rows), 'status_counts': dict(counts)}
    # leftover prev_index entries are previous SKU rows not matched this time
    if all_headers:
        for key, q in prev_index.items():
            vendor, product_key, sku_key = key
            while q:
                prev = q.popleft()
                missing_rows.append((prev.get('_values') or []) + [prev_dir.name if prev_dir else '', str(prev_dir) if prev_dir else '', '前次存在本次缺失'] + [''] * (len(TREND_COLUMNS) - 3))
        total_path = run_dir / f'SKU厂商价格趋势总表_{run_dir.name}.xlsx'
        wb = Workbook()
        ws = wb.active; ws.title = '全量明细'; ws.append(all_headers)
        for r in all_rows: ws.append(r)
        ws2 = wb.create_sheet('新产品新SKU'); ws2.append(all_headers)
        for r in new_rows: ws2.append(r)
        ws3 = wb.create_sheet('本次缺失SKU'); ws3.append(all_headers)
        for r in missing_rows: ws3.append(r)
        ws4 = wb.create_sheet('摘要'); ws4.append(['厂商','结果行数','匹配成功','新产品','新SKU','无前次基准'])
        for vendor, info in summary['vendors'].items():
            c = info['status_counts']
            ws4.append([vendor, info['rows'], c.get('匹配成功',0), c.get('新产品',0), c.get('新SKU',0), c.get('无前次基准',0)])
        wb.save(total_path)
        summary['total_trend_file'] = str(total_path)
        summary['missing_sku_rows'] = len(missing_rows)
        summary['new_rows'] = len(new_rows)
    report_path = run_dir / '趋势对比摘要.json'
    report_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    zip_path = run_dir / f'Trip_SKU厂商价格趋势_{run_dir.name}.zip'
    if zip_path.exists(): zip_path.unlink()
    exclude_from_package = {'trip-score-cdp-help.txt', 'finalize_output.json'}
    with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        for pattern in ('*.xlsx','*.md','*.json','*.txt'):
            for p in sorted(run_dir.glob(pattern)):
                if p.name == zip_path.name or p.name.startswith('.') or p.name in exclude_from_package:
                    continue
                z.write(p, arcname=p.name)
    summary['package'] = str(zip_path)
    report_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    if len(sys.argv) != 2:
        raise SystemExit('usage: trip_sku_trend_finalize.py RUN_DIR')
    main(sys.argv[1])
