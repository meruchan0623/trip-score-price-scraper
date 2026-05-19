from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def export_rows(
    rows: list[dict],
    output_path: str,
    *,
    rating_only: bool,
    extract_real_price: bool,
    extract_product_info: bool,
    include_vendor_column: bool = False,
) -> str:
    path = Path(output_path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Trip评分数据" if rating_only else "Trip商品数据"

    max_sku_layers = 0
    for row in rows:
        for key in row:
            if key.startswith("SKU层级"):
                try:
                    max_sku_layers = max(max_sku_layers, int(key.replace("SKU层级", "")))
                except ValueError:
                    pass

    fields: list[str] = []
    if include_vendor_column:
        fields.append("厂商")
    fields.extend(["商品ID", "URL", "商品名称"])
    if not rating_only:
        for index in range(1, max(max_sku_layers, 1) + 1):
            fields.append(f"SKU层级{index}")
        fields.extend(["货币", "外显价格"])
        if extract_real_price:
            fields.append("实价")
        fields.append("促销标签")
    fields.extend(["预定数", "评价数", "评分", "客户服务", "方便程度", "讯号状态", "中差评价数"])
    if extract_product_info:
        fields.append("商品资讯")

    sheet.append(fields)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append([row.get(field, "") for field in fields])

    for column_cells in sheet.columns:
        width = 12
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(len(value) + 2, 60))
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    workbook.save(path)
    return str(path)
