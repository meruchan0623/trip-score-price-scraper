import json
import os
import threading
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace
from urllib.error import HTTPError

import openpyxl
import pytest

from scripts import trip_outerlist_mode2_api_daily as collector


def test_parse_config_preserves_legacy_defaults():
    config = collector.parse_config([])

    assert config.workers == 5
    assert config.sites == ('hk',)
    assert config.mode2_fallback == 'auto'
    assert config.mode2_workers == 5
    assert config.mode2_workers_explicit is False
    assert config.api_retries == 0
    assert config.api_retry_interval_sec == 0.0
    assert config.resume_run_dir is None
    assert config.failed_only is False


def test_parse_config_accepts_valid_policy_values(tmp_path):
    resume_dir = tmp_path / 'prior-run'
    config = collector.parse_config([
        '--workers', '8',
        '--sites', 'hk,tw',
        '--output-root', str(tmp_path),
        '--baseline-workbook', 'baseline.xlsx',
        '--run-suffix', 'manual',
        '--mode2-fallback', 'off',
        '--mode2-workers', '3',
        '--api-retries', '2',
        '--api-retry-interval-sec', '0.25',
        '--resume-run-dir', str(resume_dir),
        '--failed-only',
    ])

    assert config.workers == 8
    assert config.sites == ('hk', 'tw')
    assert config.output_root == tmp_path
    assert config.baseline_workbook == Path('baseline.xlsx')
    assert config.run_suffix == 'manual'
    assert config.mode2_fallback == 'off'
    assert config.mode2_workers == 3
    assert config.mode2_workers_explicit is True
    assert config.api_retries == 2
    assert config.api_retry_interval_sec == 0.25
    assert config.resume_run_dir == resume_dir
    assert config.failed_only is True


@pytest.mark.parametrize(
    ('option', 'value'),
    [
        ('--workers', '0'),
        ('--workers', '-1'),
        ('--mode2-workers', '0'),
        ('--mode2-workers', '-1'),
        ('--mode2-workers', '6'),
        ('--api-retries', '-1'),
        ('--api-retry-interval-sec', '-0.1'),
        ('--api-retry-interval-sec', 'nan'),
        ('--api-retry-interval-sec', 'inf'),
    ],
)
def test_parse_config_rejects_invalid_numeric_values(option, value):
    with pytest.raises(SystemExit):
        collector.parse_config([option, value])


def test_parse_config_rejects_invalid_mode2_policy():
    with pytest.raises(SystemExit):
        collector.parse_config(['--mode2-fallback', 'sometimes'])


def test_runtime_policy_accepts_mode2_and_still_rejects_resume_requests(tmp_path):
    required = collector.parse_config(['--mode2-fallback', 'required'])
    collector.validate_runtime_policy(required)

    explicit_workers = collector.parse_config(['--mode2-workers', '3'])
    collector.validate_runtime_policy(explicit_workers)

    resume = collector.parse_config(['--resume-run-dir', str(tmp_path), '--failed-only'])
    with pytest.raises(collector.UnsupportedPolicyError, match='resume'):
        collector.validate_runtime_policy(resume)


def test_build_run_result_is_serializable_text_compatible_and_credential_free(tmp_path):
    result = collector.build_run_result(
        status='failed',
        phase='api_collection',
        run_id='260908-mode2-outerlist-api-daily',
        outerlist_count=12,
        api_success_count=10,
        mode2_recovered_count=0,
        failed_ids=['HK/123', 456],
        baseline_path=tmp_path / 'baseline.xlsx',
        workbook=tmp_path / 'daily.xlsx',
        errors=['Authorization: Bearer secret-token; api_key=hunter2'],
    )

    encoded = json.dumps(result)
    assert result['status'] == 'failed'
    assert result['phase'] == 'api_collection'
    assert result['run_id'] == '260908-mode2-outerlist-api-daily'
    assert result['outerlist_count'] == '12'
    assert result['api_success_count'] == '10'
    assert result['mode2_recovered_count'] == '0'
    assert result['failed_ids'] == ['HK/123', '456']
    assert result['delta_basis'] == 'twoYearSales'
    assert 'yearlySales' not in encoded
    assert result['baseline_path'] == str(tmp_path / 'baseline.xlsx')
    assert result['workbook'] == str(tmp_path / 'daily.xlsx')
    assert 'secret-token' not in encoded
    assert 'hunter2' not in encoded


@pytest.mark.parametrize('error', [
    "headers={'Authorization': 'Bearer real-secret'}",
    'https://example.invalid/path?api_key=real-secret',
    'Bearer real-secret',
    'access_token=real-secret',
    'refresh_token=real-secret',
    'client_secret=real-secret',
    'X-Amz-Credential=real-secret',
    'X-Amz-Signature=real-secret',
    'timeout while connecting',
])
def test_run_result_allows_only_reviewed_diagnostics(error):
    result = collector.build_run_result(status='failed', phase='api_collection', run_id='run', errors=[error])

    encoded = json.dumps(result)
    assert result['errors'] == ['[REDACTED]']
    assert 'real-secret' not in encoded


def test_run_result_preserves_allowlisted_diagnostics():
    result = collector.build_run_result(
        status='failed', phase='api_collection', run_id='run',
        errors=['Mode2 recovery incomplete; failed API rows remain',
                'exception_type=TimeoutError'],
    )

    assert result['errors'] == [
        'Mode2 recovery incomplete; failed API rows remain',
        'exception_type=TimeoutError',
    ]


def test_clear_completion_marker_removes_stale_marker_from_failed_terminal_path(tmp_path):
    marker = tmp_path / '.collection_done'
    marker.write_text('{"status": "success"}', encoding='utf-8')

    collector.clear_completion_marker(tmp_path)

    assert not marker.exists()


def test_previous_metrics_ignores_failed_completed_marker_but_accepts_legacy(tmp_path, monkeypatch):
    target_day = collector.date(2026, 9, 8)
    legacy = tmp_path / '260907-legacy'
    failed = tmp_path / '260907-failed'
    for run in (legacy, failed):
        run.mkdir()
        (run / '.collection_done').write_text('{}', encoding='utf-8')
        (run / '260907-0600_Trip_DailySalesVolume_FromOuterlist.xlsx').write_bytes(b'placeholder')
    (failed / 'run_result.json').write_text(json.dumps({'status': 'failed'}), encoding='utf-8')
    monkeypatch.setattr(collector, 'metrics_from_book', lambda book: {'selected': book.parent.name})

    metrics, selected_run = collector.previous_metrics(tmp_path, target_day)

    assert selected_run == legacy
    assert metrics == {'selected': '260907-legacy'}


def test_previous_metrics_ignores_newer_other_job_run_without_sales_workbook(tmp_path):
    target_day = collector.date(2026, 9, 21)
    sales_run = tmp_path / '260920-mode2-outerlist-api-daily-cron'
    sales_book = sales_run / '260920-0600_Trip_DailySalesVolume_FromOuterlist.xlsx'
    decoy = tmp_path / '260920-mode2-daily'
    _write_baseline_book(sales_book, 100)
    for run in (sales_run, decoy):
        run.mkdir(exist_ok=True)
        (run / '.collection_done').write_text('{}', encoding='utf-8')
        (run / 'run_result.json').write_text(json.dumps({'status': 'success'}), encoding='utf-8')
    os.utime(decoy / '.collection_done', (4102444800, 4102444800))

    metrics, selected_run = collector.previous_metrics(tmp_path, target_day)

    assert selected_run == sales_run
    assert metrics[('HK', '1')]['total'] == 100
    output = tmp_path / 'out.xlsx'
    collector.create_outerlist_book(output, target_day, tmp_path, [])
    collector.add_sales_sheets(output, target_day, {'HK': [{'商品ID': '1', '总销量': '150'}]}, metrics,
                               collector.baseline_day(selected_run))
    assert _header_row(output, 'HK')[6] == '销量增量（0921-0920）'


@pytest.mark.parametrize('status', ('failed', None))
def test_previous_metrics_rejects_incomplete_or_failed_sales_run_despite_decoy(tmp_path, status):
    target_day = collector.date(2026, 9, 21)
    sales_run = tmp_path / '260920-mode2-outerlist-api-daily-cron'
    _write_baseline_book(sales_run / '260920-0600_Trip_DailySalesVolume_FromOuterlist.xlsx', 100)
    (sales_run / '.collection_done').write_text('{}', encoding='utf-8')
    if status:
        (sales_run / 'run_result.json').write_text(json.dumps({'status': status}), encoding='utf-8')
    elif (sales_run / '.collection_done').exists():
        (sales_run / '.collection_done').unlink()
    decoy = tmp_path / '260920-mode2-daily'
    decoy.mkdir()
    (decoy / '.collection_done').write_text('{}', encoding='utf-8')
    (decoy / 'run_result.json').write_text(json.dumps({'status': 'success'}), encoding='utf-8')

    assert collector.previous_metrics(tmp_path, target_day) == ({}, None)


def test_fetch_keeps_yearly_sales_separate_from_two_year_delta_basis(monkeypatch):
    response_data = {
        'ResponseStatus': {'Ack': 'Success'},
        'productInfos': [{
            'productId': '123',
            'salesVolume': {
                'monthlySales': 7,
                'yearlySales': 99,
                'twoYearSales': 200,
            },
            'commentInfo': {},
            'basicInfo': {},
        }],
    }

    class Response:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    @contextmanager
    def fake_urlopen(_request, timeout):
        assert timeout == 35
        response = Response()
        response.read = lambda: json.dumps(response_data).encode()
        yield response

    monkeypatch.setattr(collector.urllib.request, 'urlopen', fake_urlopen)
    row = collector.fetch('hk', 'zh-HK', 'HKD', {'id': '123', 'vendor': 'vendor'})

    assert row['年销量'] == '99'
    assert row['总销量'] == '200'


def test_fetch_sends_frontend_consistent_headers_for_requested_hk_site_and_locale(monkeypatch):
    captured = []

    @contextmanager
    def fake_urlopen(request, timeout):
        captured.append(request)
        assert timeout == 35
        yield SimpleNamespace(read=lambda: json.dumps({
            'ResponseStatus': {'Ack': 'Success'},
            'productInfos': [{'productId': '123', 'salesVolume': {}}],
        }).encode())

    monkeypatch.setattr(collector.urllib.request, 'urlopen', fake_urlopen)

    collector.fetch('hk', 'zh-HK', 'HKD', {'id': '123', 'vendor': 'vendor'})

    headers = {name.lower(): value for name, value in captured[0].header_items()}
    assert headers['user-agent'] == collector.FRONTEND_USER_AGENT
    assert headers['accept'] == 'application/json, text/plain, */*'
    assert 'zh-HK' in headers['accept-language']
    assert headers['origin'] == 'https://hk.trip.com'
    assert headers['referer'] == 'https://hk.trip.com/things-to-do/detail/123/'
    assert headers['sec-fetch-site'] == 'same-origin'
    assert headers['sec-fetch-mode'] == 'cors'
    assert headers['sec-fetch-dest'] == 'empty'


def test_fetch_derives_frontend_headers_for_non_hk_site_and_locale(monkeypatch):
    captured = []

    @contextmanager
    def fake_urlopen(request, timeout):
        captured.append(request)
        yield SimpleNamespace(read=lambda: json.dumps({
            'ResponseStatus': {'Ack': 'Success'},
            'productInfos': [{'productId': '456', 'salesVolume': {}}],
        }).encode())

    monkeypatch.setattr(collector.urllib.request, 'urlopen', fake_urlopen)

    collector.fetch('ru', 'ru-RU', 'RUB', {'id': '456', 'vendor': 'vendor'})

    headers = {name.lower(): value for name, value in captured[0].header_items()}
    assert headers['user-agent'] == collector.FRONTEND_USER_AGENT
    assert 'ru-RU' in headers['accept-language']
    assert headers['origin'] == 'https://ru.trip.com'
    assert headers['referer'] == 'https://ru.trip.com/things-to-do/detail/456/'


def test_fetch_retries_http_430_risk_block_with_bounded_jitter(monkeypatch):
    attempts, delays = [], []

    @contextmanager
    def success_response():
        yield SimpleNamespace(read=lambda: json.dumps({
            'ResponseStatus': {'Ack': 'Success'},
            'productInfos': [{'productId': '123', 'salesVolume': {}}],
        }).encode())

    def fake_urlopen(_request, timeout):
        attempts.append(timeout)
        if len(attempts) == 1:
            raise HTTPError('https://hk.trip.com', 430, 'blocked', {}, None)
        return success_response()

    monkeypatch.setattr(collector.urllib.request, 'urlopen', fake_urlopen)
    monkeypatch.setattr(collector.random, 'uniform', lambda _low, _high: 0.125)
    monkeypatch.setattr(collector.time, 'sleep', delays.append)

    row = collector.fetch('hk', 'zh-HK', 'HKD', {'id': '123', 'vendor': 'vendor'}, api_retries=0)

    assert len(attempts) == 2
    assert delays == [1.625]
    assert row['商品ID'] == '123'
    assert not row.get('抓取备注')


def test_api_request_pacer_spaces_concurrent_requests(monkeypatch):
    clock = {'value': 10.0}
    sleeps = []

    monkeypatch.setattr(collector.time, 'monotonic', lambda: clock['value'])

    def fake_sleep(delay):
        sleeps.append(delay)
        clock['value'] += delay

    monkeypatch.setattr(collector.time, 'sleep', fake_sleep)
    monkeypatch.setattr(collector.random, 'uniform', lambda _low, _high: 0.0)
    pacer = collector.ApiRequestPacer(workers=4)
    barrier = threading.Barrier(4)

    def paced_request():
        barrier.wait()
        pacer.wait()

    threads = [threading.Thread(target=paced_request) for _ in range(4)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    assert sleeps == [collector.API_REQUEST_INTERVAL_PER_WORKER_SEC * 4] * 3


def _run_offline_main(monkeypatch, tmp_path, api_rows, mode2_rows, *, policy, workers=5, previous=None):
    product_ids = list(api_rows)
    source = tmp_path / 'outerlist-source'
    outer_rows = [
        {
            '厂商': 'Vendor',
            '商品ID': product_id,
            'URL': f'https://hk.trip.com/things-to-do/detail/{product_id}/',
            '商品名称': f'Outer {product_id}',
        }
        for product_id in product_ids
    ]
    monkeypatch.setattr(collector, 'latest_outerlist', lambda: (source, [{}]))
    monkeypatch.setattr(collector, 'outerlist_rows', lambda _items: outer_rows)
    monkeypatch.setattr(collector, 'previous_metrics', lambda _root, _day: (previous or {}, None))

    def fake_fetch(_site, _locale, _currency, item, **_kwargs):
        row = {
            '厂商': item['vendor'],
            '商品ID': item['id'],
            'URL': f'https://hk.trip.com/things-to-do/detail/{item["id"]}/',
        }
        row.update(api_rows[item['id']])
        return row

    monkeypatch.setattr(collector, 'fetch', fake_fetch)
    commands = []

    def fake_mode2_command(command, _timeout_sec):
        command = list(command)
        commands.append(command)
        if 'build' in command:
            binary = Path(command[command.index('-o') + 1])
            binary.parent.mkdir(parents=True, exist_ok=True)
            binary.touch()
        if '-output' in command:
            output = Path(command[command.index('-output') + 1])
            output.parent.mkdir(parents=True, exist_ok=True)
            output.touch()
        return 0

    monkeypatch.setattr(collector, '_run_mode2_command', fake_mode2_command)

    def fake_mode2_read(output):
        urls = Path(output).with_suffix('.urls.txt').read_text(encoding='utf-8').splitlines()
        ids = [url.rstrip('/').rsplit('/', 1)[-1] for url in urls]
        return {product_id: dict(mode2_rows[product_id]) for product_id in ids if product_id in mode2_rows}

    monkeypatch.setattr(collector, 'read_mode2_rows', fake_mode2_read, raising=False)
    args = [
        '--output-root', str(tmp_path),
        '--mode2-fallback', policy,
        '--mode2-workers', str(workers),
    ]
    return commands, args


def _only_run(tmp_path):
    return next(tmp_path.glob('*-mode2-outerlist-api-daily'))


def _sheet_rows(workbook, sheet):
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ws = wb[sheet]
    headers = ['销量增量' if str(cell.value).startswith('销量增量') else cell.value for cell in ws[2]]
    rows = {
        str(values[headers.index('商品ID')]): {
            header: '' if values[index] is None else str(values[index])
            for index, header in enumerate(headers)
        }
        for values in ws.iter_rows(min_row=3, values_only=True)
    }
    wb.close()
    return rows


def _header_row(workbook, sheet):
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    headers = [cell.value for cell in wb[sheet][2]]
    wb.close()
    return headers


def _write_baseline_book(path, total):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'HK'
    ws.append(['抓取日期'])
    ws.append(['商品ID', '总销量', '总分'])
    ws.append(['1', total, 4.5])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_sales_delta_header_uses_current_and_explicit_baseline_dates(tmp_path):
    day = collector.date(2026, 9, 14)
    baseline = tmp_path / '260913-mode2-outerlist-api-daily' / '260913-0800_Trip_DailySalesVolume_FromOuterlist.xlsx'
    _write_baseline_book(baseline, 100)
    output = tmp_path / 'out.xlsx'
    collector.create_outerlist_book(output, day, tmp_path, [])
    rows = {'HK': [{'厂商': 'V', '商品ID': '1', '总销量': '150', '总分': '4.7'}]}

    collector.add_sales_sheets(output, day, rows, collector.metrics_from_book(baseline), collector.baseline_day(baseline))

    headers = _header_row(output, 'HK')
    assert headers[6] == '销量增量（0914-0913）'
    assert headers[:6] == collector.HEADERS[:6] and headers[7:] == collector.HEADERS[7:]
    assert _sheet_rows(output, 'HK')['1']['销量增量'] == '50'


def test_sales_delta_header_reflects_explicit_baseline_from_a_different_day(tmp_path):
    day = collector.date(2026, 9, 14)
    baseline = tmp_path / '260901-1200_Trip_DailySalesVolume_FromOuterlist.xlsx'
    _write_baseline_book(baseline, 100)
    output = tmp_path / 'out.xlsx'
    collector.create_outerlist_book(output, day, tmp_path, [])

    collector.add_sales_sheets(output, day, {'HK': [{'商品ID': '1', '总销量': '120'}]}, collector.metrics_from_book(baseline), collector.baseline_day(baseline))

    assert _header_row(output, 'HK')[6] == '销量增量（0914-0901）'
    assert collector.baseline_day(tmp_path / '260913-mode2-outerlist-api-daily') == collector.date(2026, 9, 13)


def test_sales_delta_header_without_baseline_does_not_invent_a_date(tmp_path):
    day = collector.date(2026, 9, 14)
    output = tmp_path / 'out.xlsx'
    collector.create_outerlist_book(output, day, tmp_path, [])

    collector.add_sales_sheets(output, day, {'HK': [{'商品ID': '1', '总销量': '120'}]}, {}, collector.baseline_day(None))

    headers = _header_row(output, 'HK')
    assert headers[6] == '销量增量（无对比基线）'
    assert '0913' not in headers[6]
    assert len(headers) == len(collector.HEADERS)
    assert _sheet_rows(output, 'HK')['1']['销量增量'] == ''


def test_mode2_off_never_invokes_go_and_keeps_failed_run_incomplete(tmp_path, monkeypatch):
    commands, args = _run_offline_main(
        monkeypatch,
        tmp_path,
        {'1': {'抓取备注': 'api failed'}},
        {},
        policy='off',
    )

    with pytest.raises(SystemExit, match='部分失败'):
        collector.main(args)

    run = _only_run(tmp_path)
    result = json.loads((run / 'run_result.json').read_text(encoding='utf-8'))
    assert commands == []
    assert result['status'] == 'failed'
    assert result['mode2_recovered_count'] == '0'
    assert result['failed_ids'] == ['HK/1']
    assert not (run / '.collection_done').exists()


def test_mode2_auto_partitions_only_failed_ids_and_merges_without_zeroing_sales(tmp_path, monkeypatch):
    api_rows = {
        '1': {'商品名称': 'API success', '总销量': '100'},
        '2': {'商品名称': 'API old name', '总销量': '250', '抓取备注': 'api failed'},
        '3': {'抓取备注': 'api failed'},
        '4': {'抓取备注': 'api failed'},
    }
    mode2_rows = {
        '2': {
            '商品ID': '2',
            'URL': 'https://hk.trip.com/mode2/detail/2/',
            '商品名称': 'Mode2 name',
            '月销量': '25',
            '年销量': '250',
            '总销量': '250',
            '评分': '4.7',
            '抓取备注': '',
        },
        '3': {'商品ID': '3', '商品名称': 'Recovered no sales', '总销量': '', '抓取备注': ''},
        '4': {'商品ID': '4', '月销量': '40', '年销量': '400', '总销量': '400', '抓取备注': ''},
    }
    previous = {
        ('HK', '2'): {'total': 200, 'score': None},
        ('HK', '3'): {'total': 10, 'score': None},
    }
    commands, args = _run_offline_main(
        monkeypatch,
        tmp_path,
        api_rows,
        mode2_rows,
        policy='auto',
        workers=2,
        previous=previous,
    )

    with pytest.raises(SystemExit, match='部分失败'):
        collector.main(args)

    run = _only_run(tmp_path)
    result = json.loads((run / 'run_result.json').read_text(encoding='utf-8'))
    worker_commands = [command for command in commands if '-run-mode' in command]
    sent_urls = []
    for command in worker_commands:
        assert command[command.index('-run-mode') + 1] == 'mode2'
        assert command[command.index('-delay-profile') + 1] == 'normal'
        assert command[command.index('-workers') + 1] == '1'
        assert '-no-headless' in command
        urls_file = Path(command[command.index('-urls-file') + 1])
        sent_urls.extend(urls_file.read_text(encoding='utf-8').splitlines())
    assert commands[0][1:] == ['build', '-o', commands[0][3], './cmd/trip-score-cdp']
    assert len(worker_commands) == 2
    assert sorted(sent_urls) == [
        'https://hk.trip.com/things-to-do/detail/2/',
        'https://hk.trip.com/things-to-do/detail/3/',
        'https://hk.trip.com/things-to-do/detail/4/',
    ]
    assert result['status'] == 'failed'
    assert result['api_success_count'] == '1'
    assert result['mode2_recovered_count'] == '2'
    assert result['failed_ids'] == ['HK/3']
    assert not (run / '.collection_done').exists()

    rows = _sheet_rows(Path(result['workbook']), 'HK')
    assert rows['2']['URL'] == 'https://hk.trip.com/mode2/detail/2/'
    assert rows['2']['商品名称'] == 'Mode2 name'
    assert rows['2']['总销量'] == '250'
    assert rows['2']['销量增量'] == '50'
    assert rows['2']['总分'] == '4.7'
    assert rows['3']['总销量'] == ''
    assert rows['3']['销量增量'] == ''
    assert rows['3']['抓取备注'] == 'api failed'
    assert rows['4']['总销量'] == '400'
    note = (run / '任务说明.md').read_text(encoding='utf-8')
    assert 'API-only 全部成功时不启动浏览器' in note
    assert 'API 失败且尝试恢复时，可能启动 headed Go Mode2' in note


def test_mode2_required_records_partial_recovery_as_non_success_without_marker(tmp_path, monkeypatch):
    api_rows = {
        '10': {'抓取备注': 'api failed'},
        '20': {'抓取备注': 'api failed'},
    }
    mode2_rows = {
        '10': {'商品ID': '10', '月销量': '50', '年销量': '500', '总销量': '500', '抓取备注': ''},
        '20': {'商品ID': '20', '总销量': '600', '抓取备注': 'mode2 failed'},
    }
    commands, args = _run_offline_main(
        monkeypatch,
        tmp_path,
        api_rows,
        mode2_rows,
        policy='required',
        workers=5,
    )

    with pytest.raises(SystemExit, match='部分失败'):
        collector.main(args)

    run = _only_run(tmp_path)
    result = json.loads((run / 'run_result.json').read_text(encoding='utf-8'))
    assert any('-run-mode' in command for command in commands)
    assert result['status'] == 'failed'
    assert result['phase'] == 'mode2_recovery'
    assert result['mode2_recovered_count'] == '1'
    assert result['failed_ids'] == ['HK/20']
    assert result['errors'] == ['Mode2 recovery incomplete; failed API rows remain']
    assert not (run / '.collection_done').exists()



def test_run_mode2_command_timeout_terminates_then_kills_process_group(monkeypatch):
    waits = []
    group_signals = []
    popen_calls = []

    class FakeProcess:
        pid = 4242

        def wait(self, timeout):
            waits.append(timeout)
            raise collector.subprocess.TimeoutExpired(cmd=['mode2'], timeout=timeout)

        def terminate(self):
            raise AssertionError('process-group TERM should be used')

        def kill(self):
            raise AssertionError('process-group KILL should be used')

    process = FakeProcess()

    def fake_popen(command, **kwargs):
        popen_calls.append((list(command), kwargs))
        return process

    monkeypatch.setattr(collector.subprocess, 'Popen', fake_popen)
    monkeypatch.setattr(
        collector.os,
        'killpg',
        lambda pid, signal_number: group_signals.append((pid, signal_number)),
    )

    assert collector._run_mode2_command(['fake-mode2'], 17) is None
    assert waits == [
        17,
        collector.MODE2_TERMINATE_GRACE_SEC,
        collector.MODE2_TERMINATE_GRACE_SEC,
    ]
    assert group_signals == [
        (4242, collector.signal.SIGTERM),
        (4242, collector.signal.SIGKILL),
    ]
    assert popen_calls == [(['fake-mode2'], {
        'cwd': collector.REPO_ROOT,
        'stdin': collector.subprocess.DEVNULL,
        'stdout': collector.subprocess.DEVNULL,
        'stderr': collector.subprocess.DEVNULL,
        'start_new_session': True,
    })]


def test_mode2_recovery_reports_go_build_failure_without_subprocess(tmp_path, monkeypatch):
    calls = []

    def fake_command(command, timeout_sec):
        calls.append((list(command), timeout_sec))
        return 2

    monkeypatch.setattr(collector, '_run_mode2_command', fake_command)

    recovered, errors = collector.run_mode2_recovery(tmp_path, ['1'], workers=1)

    assert recovered == {}
    assert errors == ['Mode2 recovery build failed']
    assert calls[0][1] == collector.MODE2_BUILD_TIMEOUT_SEC
    assert all('-run-mode' not in command for command, _timeout in calls)


def test_mode2_recovery_reports_nonzero_worker_without_subprocess(tmp_path, monkeypatch):
    timeouts = []

    def fake_command(command, timeout_sec):
        timeouts.append(timeout_sec)
        return 0 if 'build' in command else 9

    monkeypatch.setattr(collector, '_run_mode2_command', fake_command)

    recovered, errors = collector.run_mode2_recovery(tmp_path, ['1'], workers=1)

    assert recovered == {}
    assert errors == ['Mode2 recovery worker failed', 'Mode2 recovery output missing']
    assert timeouts == [
        collector.MODE2_BUILD_TIMEOUT_SEC,
        collector.MODE2_WORKER_TIMEOUT_SEC,
    ]


def test_mode2_recovery_merges_complete_partial_workbook_after_nonzero_exit(tmp_path, monkeypatch):
    def fake_command(command, _timeout):
        if '-output' in command:
            Path(command[command.index('-output') + 1]).touch()
            return 4
        return 0

    monkeypatch.setattr(collector, '_run_mode2_command', fake_command)
    monkeypatch.setattr(collector, 'read_mode2_rows', lambda _output: {
        '1': {'商品ID': '1', '月销量': '1', '年销量': '10', '总销量': '20', '抓取备注': ''},
    })

    recovered, errors = collector.run_mode2_recovery(tmp_path, ['1'], workers=1)

    assert recovered['1']['总销量'] == '20'
    assert errors == []


def test_mode2_recovery_runs_sequential_batches_of_twenty_per_worker(tmp_path, monkeypatch):
    commands = []

    def fake_command(command, _timeout):
        commands.append(list(command))
        if '-output' in command:
            Path(command[command.index('-output') + 1]).touch()
        return 0

    monkeypatch.setattr(collector, '_run_mode2_command', fake_command)
    monkeypatch.setattr(collector, 'read_mode2_rows', lambda _output: {})

    collector.run_mode2_recovery(tmp_path, [str(value) for value in range(41)], workers=1)

    worker_commands = [command for command in commands if '-run-mode' in command]
    assert len(worker_commands) == 3
    assert [len(Path(command[command.index('-urls-file') + 1]).read_text(encoding='utf-8').splitlines())
            for command in worker_commands] == [20, 20, 1]


def test_mode2_recovery_reports_missing_worker_output_without_subprocess(tmp_path, monkeypatch):
    monkeypatch.setattr(collector, '_run_mode2_command', lambda _command, _timeout: 0)

    recovered, errors = collector.run_mode2_recovery(tmp_path, ['1'], workers=1)

    assert recovered == {}
    assert errors == ['Mode2 recovery output missing']


def test_mode2_recovery_reports_malformed_workbook_without_subprocess(tmp_path, monkeypatch):
    def fake_command(command, _timeout):
        if '-output' in command:
            output = Path(command[command.index('-output') + 1])
            output.write_bytes(b'not an xlsx workbook')
        return 0

    monkeypatch.setattr(collector, '_run_mode2_command', fake_command)

    recovered, errors = collector.run_mode2_recovery(tmp_path, ['1'], workers=1)

    assert recovered == {}
    assert errors == ['Mode2 recovery workbook invalid']


def test_mode2_worker_timeout_leaves_credential_free_failed_run_result(tmp_path, monkeypatch):
    _commands, args = _run_offline_main(
        monkeypatch,
        tmp_path,
        {'1': {'抓取备注': 'api failed'}},
        {},
        policy='auto',
    )

    def fake_timeout(command, timeout_sec):
        if 'build' in command:
            assert timeout_sec == collector.MODE2_BUILD_TIMEOUT_SEC
            return 0
        assert timeout_sec == collector.MODE2_WORKER_TIMEOUT_SEC
        return None

    monkeypatch.setattr(collector, '_run_mode2_command', fake_timeout)

    with pytest.raises(SystemExit, match='部分失败'):
        collector.main(args)

    run = _only_run(tmp_path)
    result_text = (run / 'run_result.json').read_text(encoding='utf-8')
    result = json.loads(result_text)
    assert result['status'] == 'failed'
    assert result['phase'] == 'mode2_recovery'
    assert result['failed_ids'] == ['HK/1']
    assert result['errors'] == [
        'Mode2 recovery worker failed',
        'Mode2 recovery output missing',
        'Mode2 recovery incomplete; failed API rows remain',
    ]
    assert 'api failed' not in result_text
    assert not (run / '.collection_done').exists()


def test_empty_id_only_mode2_payload_does_not_clear_api_failure(tmp_path, monkeypatch):
    _commands, args = _run_offline_main(
        monkeypatch,
        tmp_path,
        {
            '1': {
                '商品名称': 'Preserve API name',
                '总销量': '250',
                '抓取备注': 'api failed',
            },
        },
        {'1': {'商品ID': '1', '抓取备注': ''}},
        policy='auto',
    )

    with pytest.raises(SystemExit, match='部分失败'):
        collector.main(args)

    run = _only_run(tmp_path)
    result = json.loads((run / 'run_result.json').read_text(encoding='utf-8'))
    assert result['status'] == 'failed'
    assert result['mode2_recovered_count'] == '0'
    assert result['failed_ids'] == ['HK/1']
    assert result['errors'] == ['Mode2 recovery incomplete; failed API rows remain']
    assert not (run / '.collection_done').exists()

    rows = _sheet_rows(Path(result['workbook']), 'HK')
    assert rows['1']['商品名称'] == 'Preserve API name'
    assert rows['1']['总销量'] == '250'
    assert rows['1']['抓取备注'] == 'api failed'


def test_url_only_mode2_payload_does_not_clear_api_failure(tmp_path, monkeypatch):
    api_url = 'https://hk.trip.com/things-to-do/detail/1/'
    _commands, args = _run_offline_main(
        monkeypatch,
        tmp_path,
        {
            '1': {
                '商品名称': 'Preserve API name',
                '总销量': '250',
                '抓取备注': 'api failed',
            },
        },
        {'1': {'商品ID': '1', 'URL': api_url, '抓取备注': ''}},
        policy='auto',
    )

    with pytest.raises(SystemExit, match='部分失败'):
        collector.main(args)

    run = _only_run(tmp_path)
    result = json.loads((run / 'run_result.json').read_text(encoding='utf-8'))
    assert result['status'] == 'failed'
    assert result['mode2_recovered_count'] == '0'
    assert result['failed_ids'] == ['HK/1']
    assert result['errors'] == ['Mode2 recovery incomplete; failed API rows remain']
    assert not (run / '.collection_done').exists()

    rows = _sheet_rows(Path(result['workbook']), 'HK')
    assert rows['1']['URL'] == api_url
    assert rows['1']['商品名称'] == 'Preserve API name'
    assert rows['1']['总销量'] == '250'
    assert rows['1']['抓取备注'] == 'api failed'


RISK_NOTE = '接口请求失败：HTTPError: HTTP Error 430: Unknown'


def test_risk_circuit_breaker_opens_on_systemic_blocks_and_resets():
    breaker = collector.RiskCircuitBreaker()
    for _ in range(collector.RISK_CIRCUIT_MIN_CONSECUTIVE - 1):
        breaker.record(True)
    assert breaker.opened is False

    breaker.record(True)
    assert breaker.opened is True

    breaker.reset()
    assert breaker.opened is False and breaker.samples == 0


def test_risk_circuit_breaker_opens_on_block_ratio_after_min_samples():
    breaker = collector.RiskCircuitBreaker(min_consecutive=99, min_samples=10, ratio=0.6)
    for _ in range(10):
        breaker.record(True)
    # Ratio rule only matters once the sample floor is reached.
    assert breaker.opened is True


def test_risk_note_detection_and_failure_row_match_fetch_shape():
    assert collector.note_is_risk_block(RISK_NOTE) is True
    assert collector.note_is_risk_block('接口请求失败：RuntimeError: returned product ID mismatch') is False

    row = collector.failure_row('hk', 'zh-HK', 'HKD', '123', 'Vendor', 'note')
    assert row == {
        '厂商': 'Vendor',
        '商品ID': '123',
        'URL': 'https://hk.trip.com/things-to-do/detail/123/?curr=HKD&locale=zh-HK',
        '抓取备注': 'note',
    }


def test_systemic_risk_block_skips_requests_and_mode2_recovery(tmp_path, monkeypatch):
    product_ids = [str(1000 + index) for index in range(30)]
    api_rows = {product_id: {'抓取备注': RISK_NOTE} for product_id in product_ids}
    calls = []

    commands, args = _run_offline_main(
        monkeypatch, tmp_path, api_rows, {}, policy='auto', workers=1,
    )

    def counting_fetch(_site, _locale, _currency, item, **_kwargs):
        calls.append(item['id'])
        return {
            '厂商': item['vendor'],
            '商品ID': item['id'],
            'URL': f'https://hk.trip.com/things-to-do/detail/{item["id"]}/',
            '抓取备注': RISK_NOTE,
        }

    def unavailable_browser(_command, _env, _timeout):
        return 1, '', 'ego-browser unavailable'

    monkeypatch.setattr(collector, 'fetch', counting_fetch)
    monkeypatch.setattr(collector, 'browser_recovery_runner', unavailable_browser)
    monkeypatch.setattr(collector, 'API_RETRY_COOLDOWN_SEC', 0)
    monkeypatch.setattr(collector.time, 'sleep', lambda _delay: None)

    with pytest.raises(SystemExit, match='部分失败'):
        collector.main(args)

    run = _only_run(tmp_path)
    result = json.loads((run / 'run_result.json').read_text(encoding='utf-8'))
    # Requests stop once the circuit opens, so the run ends quickly instead of
    # replaying the same rejection for every product.
    assert len(calls) < len(product_ids)
    assert result['status'] == 'failed'
    assert result['failed_ids'] == ['HK/' + product_id for product_id in product_ids]
    # The browser Mode2 export has no sales volume, so recovery cannot fix a
    # risk-blocked sales report; it must not be attempted at all.
    assert [command for command in commands if '-run-mode' in command] == []
    assert 'risk_circuit_opened' in result['errors']
    assert 'browser_recovery_failed' in result['errors']
    assert 'mode2_recovery_skipped_risk_block' in result['errors']
    assert not (run / '.collection_done').exists()


def test_browser_recovery_fills_blocked_rows_and_completes(tmp_path, monkeypatch):
    """A browser-session recovery of every blocked row makes the day succeed."""
    product_ids = [str(2000 + index) for index in range(3)]
    api_rows = {product_id: {'抓取备注': RISK_NOTE} for product_id in product_ids}
    commands, args = _run_offline_main(
        monkeypatch, tmp_path, api_rows, {}, policy='off', workers=1,
    )

    def blocked_fetch(_site, _locale, _currency, item, **_kwargs):
        return {
            '厂商': item['vendor'],
            '商品ID': item['id'],
            'URL': f'https://hk.trip.com/things-to-do/detail/{item["id"]}/',
            '抓取备注': RISK_NOTE,
        }

    def fake_runner(_command, env, _timeout):
        ids = json.loads(Path(env['TRIP_RECOVERY_IDS']).read_text(encoding='utf-8'))
        lines = []
        for product_id in ids:
            lines.append(json.dumps({'id': product_id, 'data': {
                'ResponseStatus': {'Ack': 'Success'},
                'productInfos': [{
                    'productId': int(product_id),
                    'basicInfo': {'productName': 'recovered ' + product_id},
                    'salesVolume': {'monthlySales': 11, 'yearlySales': 22, 'twoYearSales': 33},
                    'commentInfo': {'commentCount': 44, 'commentScore': 4.5},
                }],
            }}))
        Path(env['TRIP_RECOVERY_OUT']).write_text('\n'.join(lines) + '\n', encoding='utf-8')
        return 0, '{"ok":3,"failed":0}', ''

    monkeypatch.setattr(collector, 'fetch', blocked_fetch)
    monkeypatch.setattr(collector, 'browser_recovery_runner', fake_runner)
    monkeypatch.setattr(collector, 'API_RETRY_COOLDOWN_SEC', 0)
    monkeypatch.setattr(collector.time, 'sleep', lambda _delay: None)

    collector.main(args)

    run = _only_run(tmp_path)
    result = json.loads((run / 'run_result.json').read_text(encoding='utf-8'))
    assert result['status'] == 'success'
    assert 'browser_recovery_applied' in result['errors']
    assert result['failed_ids'] == []

    workbook = openpyxl.load_workbook(result['workbook'], data_only=True)
    sheet = workbook['HK']
    header_row = 0
    header = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if '商品ID' in [str(value) for value in row if value is not None]:
            header_row, header = index, [str(value) for value in row]
            break
    assert header_row, 'HK sheet header row not found'
    rows = [
        dict(zip(header, [str(cell.value) if cell.value is not None else '' for cell in row]))
        for row in sheet.iter_rows(min_row=header_row + 1)
        if row[0].value
    ]
    recovered = {str(row['商品ID']): row for row in rows}
    assert set(recovered) == set(product_ids)
    for row in recovered.values():
        assert str(row['总销量']) == '33'
        assert str(row['总分']) == '4.5'


def test_cool_down_retry_recovers_products_after_transient_block(tmp_path, monkeypatch):
    product_ids = [str(2000 + index) for index in range(14)]
    api_rows = {product_id: {'抓取备注': RISK_NOTE} for product_id in product_ids}
    blocked = set(product_ids)
    calls = []

    commands, args = _run_offline_main(
        monkeypatch, tmp_path, api_rows, {}, policy='off', workers=1,
    )

    def flaky_fetch(_site, _locale, _currency, item, **_kwargs):
        calls.append(item['id'])
        if len(calls) > collector.RISK_CIRCUIT_MIN_CONSECUTIVE:
            # The first pass is blocked wholesale (the circuit opens after the
            # consecutive-block threshold); the cool-down re-probe is answered.
            blocked.discard(item['id'])
        if item['id'] in blocked:
            return {
                '厂商': item['vendor'],
                '商品ID': item['id'],
                'URL': f'https://hk.trip.com/things-to-do/detail/{item["id"]}/',
                '抓取备注': RISK_NOTE,
            }
        return {
            '厂商': item['vendor'],
            '商品ID': item['id'],
            'URL': f'https://hk.trip.com/things-to-do/detail/{item["id"]}/',
            '商品名称': 'Recovered',
            '总销量': '321',
        }

    monkeypatch.setattr(collector, 'fetch', flaky_fetch)
    monkeypatch.setattr(collector, 'API_RETRY_COOLDOWN_SEC', 0)
    monkeypatch.setattr(collector.time, 'sleep', lambda _delay: None)

    # main() returns None on success (it prints the delivery summary).
    collector.main(args)

    run = _only_run(tmp_path)
    result = json.loads((run / 'run_result.json').read_text(encoding='utf-8'))
    assert result['status'] == 'success'
    assert result['api_success_count'] == str(len(product_ids))
    assert result['failed_ids'] == []
    assert len(calls) > len(product_ids)
    rows = _sheet_rows(Path(result['workbook']), 'HK')
    assert rows[product_ids[0]]['总销量'] == '321'
    assert (run / '.collection_done').exists()


def test_browser_recovery_replacements_parse_in_page_responses(tmp_path):
    """Blocked rows are rebuilt from the browser session's in-page API body."""
    payload = {
        'ResponseStatus': {'Ack': 'Success'},
        'productInfos': [{
            'productId': 57042119,
            'basicInfo': {'productName': '中國內地 5G eSIM'},
            'salesVolume': {'monthlySales': 11, 'yearlySales': 22, 'twoYearSales': 33},
            'commentInfo': {'commentCount': 44, 'commentScore': 4.5},
        }],
    }
    commands = []

    def runner(command, env, timeout):
        commands.append(command)
        Path(env['TRIP_RECOVERY_OUT']).write_text(
            json.dumps({'id': '57042119', 'data': payload}) + '\n', encoding='utf-8'
        )
        return 0, '{"ok":1,"failed":0}', ''

    blocked = [{'厂商': 'TSIM 網址', '商品ID': '57042119', '抓取备注': '接口请求失败：' + RISK_NOTE}]
    replacements, problem = collector.browser_recovery_replacements(
        blocked, 'hk', 'zh-HK', 'HKD', tmp_path, runner=runner
    )

    assert problem == ''
    assert commands and Path(commands[0][0]).name == 'ego-browser' and commands[0][1] == 'nodejs'
    row = replacements['57042119']
    assert row['总销量'] == '33' and row['年销量'] == '22' and row['总分'] == '4.5'
    assert not row.get('抓取备注')


def test_browser_recovery_reports_missing_output(tmp_path):
    """A browser run that produces no JSONL is a diagnostic, not fake data."""
    def runner(command, env, timeout):
        return 1, '', 'ego-browser exploded'

    replacements, problem = collector.browser_recovery_replacements(
        [{'厂商': 'X', '商品ID': '1'}], 'hk', 'zh-HK', 'HKD', tmp_path, runner=runner
    )

    assert replacements == {}
    assert 'ego-browser exploded' in problem
