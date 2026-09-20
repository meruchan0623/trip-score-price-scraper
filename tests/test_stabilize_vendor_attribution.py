from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).parents[1]
SCRIPT = REPO_ROOT / "scripts" / "stabilize_vendor_attribution.py"
SPEC = importlib.util.spec_from_file_location("stabilize_vendor_attribution", SCRIPT)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)

SITES = ["TW", "HK", "RU", "JP", "VN", "TH"]
WORKBOOK_HEADER = ["厂商", "商品ID", "URL", "商品名称", "总促销标签"]
# Mirrors the live Google Sheets export: numeric IDs and the `TSIM` header name.
TARGET_HEADER = ["TSIM ID", "分類", "目的地", "TSIM", "BC", "TGT", "FR", "JY", "其他"]


def _write_target(path: Path, rows: list[list]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "工作表1"
    ws.append(TARGET_HEADER)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def _write_workbook(path: Path, rows_per_site: dict[str, list[tuple[str, str, str]]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for site in SITES:
        ws = wb.create_sheet(site)
        ws.append(WORKBOOK_HEADER)
        for vendor, pid, url in rows_per_site.get(site, []):
            ws.append([vendor, pid, url, "name", ""])
    wb.save(path)
    return path


def _vendor_column(path: Path, site: str) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [list(r) for r in wb[site].iter_rows(values_only=True)]
    wb.close()
    header = rows[0]
    idx = header.index("厂商")
    return [str(r[idx]).strip() for r in rows[1:] if r]


def test_product_id_normalizes_google_export_values() -> None:
    assert MODULE.product_id(85473945.0) == "85473945"
    assert MODULE.product_id("85473945.0") == "85473945"
    assert MODULE.product_id(85473945) == "85473945"
    assert MODULE.product_id(" 85473945 ") == "85473945"
    assert MODULE.product_id("https://hk.trip.com/things-to-do/detail/85473945/?curr=HKD") == "85473945"
    assert MODULE.product_id(None) == ""
    assert MODULE.product_id("") == ""
    assert MODULE.product_id("Europe 41 Regions") == ""
    assert MODULE.product_id(85473945.5) == ""


def test_target_map_reads_float_ids_and_tsim_alias(tmp_path: Path) -> None:
    target = _write_target(
        tmp_path / "target.xlsx",
        [
            [1001.0, "歐洲", "Europe", 1001.0, 2001.0, 3001.0, 4001.0, 5001.0, None],
            [1002.0, "歐洲", "Turkey", 1002.0, None, None, None, None, None],
        ],
    )
    unique, conflicts, count = MODULE.target_map(target)
    assert count == 6
    assert conflicts == {}
    assert unique == {
        "1001": "TSIM 網址",
        "2001": "BC",
        "3001": "TGT",
        "4001": "FR",
        "5001": "JY",
        "1002": "TSIM 網址",
    }


def test_conflicts_use_manual_reference_then_join(tmp_path: Path) -> None:
    target = _write_target(
        tmp_path / "target.xlsx",
        [[9001.0, "亞洲", "Asia", None, 9001.0, 9001.0, None, None, None]],
    )
    unique, conflicts, _ = MODULE.target_map(target)
    assert unique == {}
    assert conflicts == {"9001": ["BC", "TGT"]}

    manual_file = tmp_path / "manual.json"
    manual_file.write_text(
        json.dumps({"lookup_by_product_id": {"9001": {"vendor_name": "TGT"}}}),
        encoding="utf-8",
    )
    manual = MODULE.manual_attribution_map(manual_file)
    assert manual == {"9001": "TGT"}

    resolved, origin = MODULE.resolve_vendors(unique, conflicts, manual)
    assert resolved == {"9001": "TGT"}
    assert origin["from_manual_on_conflict"] == 1

    joined, origin_joined = MODULE.resolve_vendors(unique, conflicts, {})
    assert joined == {"9001": "BC/TGT"}
    assert origin_joined["joined_conflict"] == 1


def test_manual_reference_naming_is_normalized(tmp_path: Path) -> None:
    manual_file = tmp_path / "manual.json"
    manual_file.write_text(
        json.dumps(
            {
                "lookup_by_product_id": {
                    "1": {"vendor_name": "TSim"},
                    "2": {"vendor_name": "OTH"},
                    "3": {"vendor_name": "unknown-ish"},
                }
            }
        ),
        encoding="utf-8",
    )
    assert MODULE.manual_attribution_map(manual_file) == {"1": "TSIM 網址", "2": "其他"}


def test_daily_workbook_gets_vendor_names_instead_of_manual_placeholder(tmp_path: Path) -> None:
    target = _write_target(
        tmp_path / "target.xlsx",
        [
            [1001.0, "歐洲", "Europe", 1001.0, 2001.0, 3001.0, None, None, None],
            [1002.0, "亞洲", "Asia", 1002.0, None, None, 4002.0, 5002.0, 6002.0],
        ],
    )
    workbook = _write_workbook(
        tmp_path / "merged.xlsx",
        {
            site: [
                ("手动输入", "1001", "https://hk.trip.com/things-to-do/detail/1001/"),
                ("手动输入", "2001", "https://hk.trip.com/things-to-do/detail/2001/"),
                ("手动输入", "3001", "https://hk.trip.com/things-to-do/detail/3001/"),
                ("手动输入", "4002", "https://hk.trip.com/things-to-do/detail/4002/"),
                ("手动输入", "5002", "https://hk.trip.com/things-to-do/detail/5002/"),
                ("手动输入", "6002", "https://hk.trip.com/things-to-do/detail/6002/"),
                ("JY", "1002", "https://hk.trip.com/things-to-do/detail/1002/"),
                ("手动输入", "999999", "https://hk.trip.com/things-to-do/detail/999999/"),
            ]
            for site in SITES
        },
    )
    unique, conflicts, _ = MODULE.target_map(target)
    vendors, _ = MODULE.resolve_vendors(unique, conflicts, {})
    stats = MODULE.stabilize(workbook, vendors)

    assert stats["filled"] == len(SITES) * 6
    assert stats["preserved"] == len(SITES)
    assert stats["unmatched"] == len(SITES)
    assert stats["missing_columns"] == 0

    for site in SITES:
        values = _vendor_column(workbook, site)
        assert values[0:3] == ["TSIM 網址", "BC", "TGT"]
        assert values[3:6] == ["FR", "JY", "其他"]
        assert values[6] == "JY"  # already named: never overwritten
        assert values[7] == "手动输入"  # unresolvable stays untouched, never invented


def _run_main(argv: list[str]) -> int:
    saved = sys.argv
    try:
        sys.argv = ["stabilize", *argv]
        return MODULE.main()
    finally:
        sys.argv = saved


def test_main_fails_when_target_yields_no_product_ids(tmp_path: Path, capsys) -> None:
    target = _write_target(
        tmp_path / "target.xlsx",
        [["text", "歐洲", "Europe", "no-id", "no-id", "no-id", "no-id", "no-id", "no-id"]],
    )
    workbook = _write_workbook(
        tmp_path / "merged.xlsx",
        {site: [("手动输入", "1001", "https://hk.trip.com/things-to-do/detail/1001/")] for site in SITES},
    )
    rc = _run_main(["--target-table", str(target), str(workbook)])
    assert rc == MODULE.EXIT_NO_TARGET_IDS
    assert "FATAL" in capsys.readouterr().err
    assert _vendor_column(workbook, "HK") == ["手动输入"]


def test_main_fails_when_no_row_can_be_resolved(tmp_path: Path, capsys) -> None:
    target = _write_target(
        tmp_path / "target.xlsx",
        [[1001.0, "歐洲", "Europe", 1001.0, 2001.0, 3001.0, 4001.0, 5001.0, None]],
    )
    workbook = _write_workbook(
        tmp_path / "merged.xlsx",
        {site: [("手动输入", "777777", "https://hk.trip.com/things-to-do/detail/777777/")] for site in SITES},
    )
    rc = _run_main(
        ["--target-table", str(target), "--manual-attribution", str(tmp_path / "none.json"), str(workbook)]
    )
    assert rc == MODULE.EXIT_NO_MATCHES
    assert "FATAL" in capsys.readouterr().err
    assert _vendor_column(workbook, "HK") == ["手动输入"]


def test_main_fills_every_legacy_row_end_to_end(tmp_path: Path) -> None:
    target = _write_target(
        tmp_path / "target.xlsx",
        [[1001.0, "歐洲", "Europe", 1001.0, 2001.0, 3001.0, 4001.0, 5001.0, None]],
    )
    workbook = _write_workbook(
        tmp_path / "merged.xlsx",
        {
            site: [
                ("手动输入", pid, f"https://hk.trip.com/things-to-do/detail/{pid}/")
                for pid in ("1001", "2001", "3001", "4001", "5001")
            ]
            for site in SITES
        },
    )
    rc = _run_main(
        ["--target-table", str(target), "--manual-attribution", str(tmp_path / "missing.json"), str(workbook)]
    )
    assert rc == 0
    for site in SITES:
        assert _vendor_column(workbook, site) == ["TSIM 網址", "BC", "TGT", "FR", "JY"]
