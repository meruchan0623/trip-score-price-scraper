from __future__ import annotations

import importlib.util
from pathlib import Path

import openpyxl

SCRIPT = Path(__file__).parents[1] / "scripts" / "hermes-cron" / "trip_mode2_yesterday_promo.py"
SPEC = importlib.util.spec_from_file_location("trip_mode2_yesterday_promo", SCRIPT)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)

SITES = ["TW", "HK", "RU", "JP", "VN", "TH"]
HEADER = ["商品名称", "商品ID", "厂商", "总促销标签", "价格"]


def _write(path: Path, per_site: dict[str, list[list]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for site in SITES:
        ws = wb.create_sheet(site)
        ws.append(HEADER)
        for row in per_site.get(site, []):
            ws.append(row)
    wb.save(path)
    return path


def _sheet_rows(path: Path, site: str) -> list[list]:
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = [list(r) for r in wb[site].iter_rows(values_only=True)]
    wb.close()
    return rows


def _labels(rows: list[list]) -> dict[str, tuple]:
    header = rows[0]
    pid, base, promo = header.index("商品ID"), header.index("昨日总促销标签"), header.index("总促销标签")
    return {str(r[pid]): (r[base], r[promo]) for r in rows[1:]}


YESTERDAY = {
    "TW": [["A", "100", "V", "預訂指定套票減50% | 減5%", 1], ["B", "200", "V", "指定產品減20%", 2]],
    "HK": [["A", "100", "V", "預訂指定套票減50%", 1]],
    "RU": [["A", "100", "V", "скидка 50 % | скидка 8 %", 1]],
    "JP": [["A", "100", "V", "予約で50%オフ 🎉", 1]],
    "VN": [["A", "100", "V", "Giảm 50% khi đặt trước", 1]],
    "TH": [["A", "100", "V", "ลด 50% เมื่อจอง", 1]],
}
TODAY = {
    # Reordered relative to yesterday, plus product 300 that has no baseline.
    "TW": [["B", "200", "V", "今日文案", 2], ["C", "300", "V", "新品", 3], ["A", "100", "V", "預訂減50%", 1]],
    "HK": [["A", "100", "V", "今日HK", 1]],
    "RU": [["A", "100", "V", "сегодня", 1]],
    "JP": [["A", "100", "V", "今日", 1]],
    "VN": [["A", "100", "V", "hôm nay", 1]],
    "TH": [["A", "100", "V", "วันนี้", 1]],
}


def _setup(tmp_path: Path) -> tuple[Path, Path]:
    previous = _write(tmp_path / "260910-mode2-daily" / "260910-0300_trip_mode2_daily_6site.xlsx", YESTERDAY)
    current = _write(tmp_path / "260911-mode2-daily" / "260911-0300_trip_mode2_daily_6site.xlsx", TODAY)
    return previous, current


def test_enrich_inserts_exact_previous_labels_before_current_column(tmp_path: Path) -> None:
    previous, current = _setup(tmp_path)
    before = previous.read_bytes()

    assert MODULE.main([str(current)]) == 0

    for site in SITES:
        rows = _sheet_rows(current, site)
        header = rows[0]
        assert header.count("昨日总促销标签") == 1
        assert header.index("昨日总促销标签") + 1 == header.index("总促销标签")
        # Every other column and value is preserved, in order.
        stripped = [h for h in header if h != "昨日总促销标签"]
        assert stripped == HEADER
        expected = {str(r[1]): r for r in TODAY[site]}
        for row in rows[1:]:
            without = row[:3] + row[4:]
            assert without == expected[str(row[1])]
        labels = _labels(rows)
        yesterday = {str(r[1]): r[3] for r in YESTERDAY.get(site, [])}
        for pid, (base, promo) in labels.items():
            assert base == yesterday.get(pid)  # exact Unicode text, None when absent
            assert promo == expected[pid][3]

    tw = _labels(_sheet_rows(current, "TW"))
    assert tw["100"] == ("預訂指定套票減50% | 減5%", "預訂減50%")
    assert tw["200"] == ("指定產品減20%", "今日文案")
    assert tw["300"] == (None, "新品")
    assert _labels(_sheet_rows(current, "RU"))["100"][0] == "скидка 50 % | скидка 8 %"
    assert _labels(_sheet_rows(current, "JP"))["100"][0] == "予約で50%オフ 🎉"
    assert previous.read_bytes() == before  # prior workbook untouched


def test_rerun_is_idempotent_without_duplicate_column(tmp_path: Path) -> None:
    _, current = _setup(tmp_path)
    assert MODULE.main([str(current)]) == 0
    first = {site: _sheet_rows(current, site) for site in SITES}
    assert MODULE.main([str(current)]) == 0
    second = {site: _sheet_rows(current, site) for site in SITES}
    assert first == second
    for site in SITES:
        assert second[site][0].count("昨日总促销标签") == 1


def test_missing_baseline_workbook_yields_blank_column(tmp_path: Path) -> None:
    current = _write(tmp_path / "260911-mode2-daily" / "260911-0300_trip_mode2_daily_6site.xlsx", TODAY)
    assert MODULE.main([str(current)]) == 0
    for site in SITES:
        rows = _sheet_rows(current, site)
        assert rows[0].index("昨日总促销标签") + 1 == rows[0].index("总促销标签")
        assert all(base is None for base, _ in _labels(rows).values())
        assert [p for _, p in _labels(rows).values()] == [r[3] for r in TODAY[site]]


def _font_colors(path: Path, site: str) -> dict[str, dict[str, str | None]]:
    """Per 商品ID: font colour rgb of every cell in the row, keyed by header."""
    wb = openpyxl.load_workbook(path)
    ws = wb[site]
    header = [c.value for c in ws[1]]
    pid_i = header.index("商品ID")
    out = {}
    for row in ws.iter_rows(min_row=2):
        colors = {}
        for h, cell in zip(header, row):
            color = cell.font.color if cell.font is not None else None
            # openpyxl returns descriptor text (not None) for theme/indexed/default colours;
            # only a true RGB colour counts as an explicit font colour.
            colors[h] = color.rgb if color is not None and color.type == "rgb" and isinstance(color.rgb, str) else None
        out[str(row[pid_i].value)] = colors
    wb.close()
    return out


RED = "FFFF0000"


def test_changed_current_label_is_red_but_equal_or_missing_baseline_is_not(tmp_path: Path) -> None:
    today = {site: [list(r) for r in rows] for site, rows in TODAY.items()}
    today["HK"][0][3] = YESTERDAY["HK"][0][3]  # equal to yesterday -> not red
    previous = _write(tmp_path / "260910-mode2-daily" / "260910-0300_trip_mode2_daily_6site.xlsx", YESTERDAY)
    current = _write(tmp_path / "260911-mode2-daily" / "260911-0300_trip_mode2_daily_6site.xlsx", today)
    before = previous.read_bytes()

    assert MODULE.main([str(current)]) == 0

    tw = _font_colors(current, "TW")
    assert tw["100"]["总促销标签"] == RED  # changed
    assert tw["200"]["总促销标签"] == RED  # changed
    assert tw["300"]["总促销标签"] is None  # no baseline -> not flagged
    assert _font_colors(current, "HK")["100"]["总促销标签"] is None  # equal -> not flagged
    for site in ("RU", "JP", "VN", "TH"):
        assert _font_colors(current, site)["100"]["总促销标签"] == RED
    # Only the current promo cell is touched; baseline and all other cells keep their default font.
    for site in SITES:
        for pid, colors in _font_colors(current, site).items():
            for h, color in colors.items():
                if h != "总促销标签":
                    assert color is None, (site, pid, h)
    assert previous.read_bytes() == before


def test_rerun_keeps_red_flags_without_duplicate_column_or_style_corruption(tmp_path: Path) -> None:
    _, current = _setup(tmp_path)
    assert MODULE.main([str(current)]) == 0
    first_rows = {site: _sheet_rows(current, site) for site in SITES}
    first_colors = {site: _font_colors(current, site) for site in SITES}
    assert MODULE.main([str(current)]) == 0
    assert {site: _sheet_rows(current, site) for site in SITES} == first_rows
    assert {site: _font_colors(current, site) for site in SITES} == first_colors
    tw = _font_colors(current, "TW")
    assert tw["100"]["总促销标签"] == RED and tw["300"]["总促销标签"] is None
    for site in SITES:
        assert first_rows[site][0].count("昨日总促销标签") == 1
    # A later run whose baseline now matches today's label clears the red flag.
    matching = {site: [[r[0], r[1], r[2], r[3], r[4]] for r in rows] for site, rows in TODAY.items()}
    baseline = _write(tmp_path / "match.xlsx", matching)
    assert MODULE.main([str(current), "--baseline", str(baseline)]) == 0
    assert all(c["总促销标签"] is None for c in _font_colors(current, "TW").values())


def test_previous_run_dir_crosses_month_and_year_boundaries(tmp_path: Path) -> None:
    jan1 = tmp_path / "260101-mode2-daily" / "260101-0300_trip_mode2_daily_6site.xlsx"
    assert MODULE.previous_run_dir(jan1) == tmp_path / "251231-mode2-daily"
    mar1 = tmp_path / "260301-mode2-daily" / "260301-0300_trip_mode2_daily_6site.xlsx"
    assert MODULE.previous_run_dir(mar1) == tmp_path / "260228-mode2-daily"  # 2026 is not a leap year
    assert MODULE.previous_run_dir(tmp_path / "bogus" / "x.xlsx") is None


def test_xlsx_in_prefers_timestamped_then_legacy_name(tmp_path: Path) -> None:
    run_dir = tmp_path / "260910-mode2-daily"
    run_dir.mkdir()
    assert MODULE.xlsx_in(run_dir) is None
    legacy = _write(run_dir / "trip_mode2_daily_6site.xlsx", YESTERDAY)
    assert MODULE.xlsx_in(run_dir) == legacy
    stamped = _write(run_dir / "260910-0300_trip_mode2_daily_6site.xlsx", YESTERDAY)
    assert MODULE.xlsx_in(run_dir) == stamped
    assert MODULE.xlsx_in(tmp_path / "nope") is None
