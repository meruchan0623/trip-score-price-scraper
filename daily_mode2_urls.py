#!/usr/bin/env python3
"""
daily_mode2_urls.py — 从 日常抓取目标.xlsx 生成 mode2 多站点 URL 列表

用法：
  python3 daily_mode2_urls.py [--output PATH] [--source-site hk]

产出：
  - 一行一个产品 URL（格式 `https://hk.trip.com/things-to-do/detail/<ID>/`）
  - 去重、只读 Sheet1/工作表1 列 TSIM 網址/BC/TGT/FR/JY/其他 中所有数字 ID
  - 默认输出到 /tmp/trip_mode2_daily_template.txt

此脚本只做纯数据转换，不启动浏览器、不发请求、不运行抓取。
"""
import argparse
import re
from pathlib import Path

import openpyxl

DEFAULT_INPUT = (
    Path(__file__).resolve().parent / "docs" / "日常抓取目标.xlsx"
)
DEFAULT_OUTPUT = Path("/tmp/trip_mode2_daily_template.txt")
VENDOR_COLS = ["TSIM 網址", "BC", "TGT", "FR", "JY", "其他"]
# The public source calls the canonical TSIM column `TSIM`; legacy mirrors
# call it `TSIM 網址`. Normalize only the header name, never product values.
VENDOR_COLUMN_ALIASES = {"TSIM 網址": ("TSIM 網址", "TSIM"), **{name: (name,) for name in VENDOR_COLS[1:]}}
# Google XLSX exports preserve the source's localized default worksheet name.
# Accept the canonical local name and the actual public-sheet export name.
TARGET_SHEET_NAMES = ("Sheet1", "工作表1")


def resolve_target_sheet(workbook: openpyxl.Workbook):
    for sheet_name in TARGET_SHEET_NAMES:
        if sheet_name in workbook.sheetnames:
            return workbook[sheet_name]
    raise KeyError(
        f"target workbook missing one of {list(TARGET_SHEET_NAMES)!r}; "
        f"found {workbook.sheetnames!r}"
    )


def resolve_vendor_columns(worksheet) -> dict[str, int]:
    columns: dict[str, int] = {}
    for canonical, aliases in VENDOR_COLUMN_ALIASES.items():
        for cell in worksheet[1]:
            if cell.value in aliases:
                columns[canonical] = cell.column
                break
    return columns


def parse_product_id(value) -> int | None:
    """Accept integer-like Excel cells while rejecting fractional/non-numeric IDs."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return int(text) if text.isdigit() else None


def extract_product_ids(path: Path) -> list[int]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = resolve_target_sheet(wb)
    col_map = resolve_vendor_columns(ws)
    ids: list[int] = []
    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        for col_name in VENDOR_COLS:
            c = col_map.get(col_name)
            if c is None:
                continue
            v = row[c - 1]
            pid = parse_product_id(v)
            if pid is not None:
                if pid not in seen:
                    seen.add(pid)
                    ids.append(pid)
    wb.close()
    return ids


def main(argv: list[str] | None = None) -> None:
    p = argparse.ArgumentParser(description="生成 mode2 多站点 URL 列表")
    p.add_argument("--input", default=str(DEFAULT_INPUT), help="日常抓取目标.xlsx 路径")
    p.add_argument("--output", default=str(DEFAULT_OUTPUT))
    p.add_argument("--source-site", default="hk", help="URL 来源站点 code（默认 hk）")
    args = p.parse_args(argv)
    ids = extract_product_ids(Path(args.input))
    lines = [
        f"https://{args.source_site}.trip.com/things-to-do/detail/{pid}/"
        for pid in ids
    ]
    Path(args.output).write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"{len(lines)} URLs -> {args.output}")


if __name__ == "__main__":
    main()
